"""Pre-intake analysis: find the sources in a folder and measure them — counts only.

Nothing here prints or stores a data value. Before governance has answered
`policies.egress`, the assistant may only see what `stats_only` allows, so the analysis
records file names, headers, row counts, empty counts, distinct counts, value LENGTHS,
how many values each date format parses, and whether candidate keys are unique.

The result (`.dwh/analysis.json`) feeds the defaults of the intake workbook: detected
types, candidate keys, null policies, date formats, classification hints. Every default
it produces is a proposal the owner keeps or changes in the workbook.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from . import config as C
from .project import Project, audit, now_iso

DATA_EXT = {".csv": "csv", ".tsv": "csv", ".txt": "csv", ".parquet": "parquet", ".json": "json",
            ".jsonl": "jsonl", ".ndjson": "jsonl", ".xlsx": "xlsx"}
BATCH_WORD = re.compile(r"(?i)(batch|part|chunk|load|drop|delivery|day|month|week|year|dt|date|snapshot|run)[_-]?$")
DATE_TOKEN = re.compile(r"\d{4}-\d{2}(-\d{2})?|\d{8}|\d{6}")

TS_FORMATS = ["%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%dT%H:%M:%S.%f",
              "%Y-%m-%dT%H:%M:%SZ", "%Y-%m-%d %H:%M", "%d/%m/%Y %H:%M:%S", "%m/%d/%Y %H:%M:%S",
              "%d/%m/%Y %H:%M", "%m/%d/%Y %H:%M", "%d-%m-%Y %H:%M:%S", "%Y/%m/%d %H:%M:%S"]
DATE_FORMATS = ["%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d.%m.%Y", "%Y%m%d"]
DATEISH_NAME = re.compile(r"(?i)(date|time|_at$|_ts$|timestamp|_dt$|^dt_|_on$|period|month|day)")

MONEY_NAME = re.compile(r"(?i)(price|amount|amt|value|fee|cost|revenue|total|tax|tip|fare|freight|payment|"
                        r"balance|mrr|arr|salary|discount|charge|spend|gmv|sales)")
CODE_ALWAYS = re.compile(r"(?i)(zip|postal|postcode|post_code|cep|pin_?code|phone|mobile|telephone)")
CODE_IF_RAGGED = re.compile(r"(?i)(code|prefix|account|acct|sku|ean|upc|isbn)")
ID_NAME = re.compile(r"(?i)(^id$|_id$|_key$|^key$|_code$|^code$|_no$|_number$|_uuid$|^uuid$)")
SEQ_NAME = re.compile(r"(?i)(_item_id$|_sequential$|_seq$|_sequence$|_line$|_line_no$|_no$|_num$|_index$|_version$)")

PII_RULES = [
    (re.compile(r"(?i)(e_?mail)"), "looks like an e-mail address"),
    (re.compile(r"(?i)(phone|mobile|telephone|^tel$|_tel$)"), "looks like a phone number"),
    (re.compile(r"(?i)(first_?name|last_?name|full_?name|surname|^name$|customer_name|client_name|"
                r"patient_name|person_name|contact_name)"), "looks like a person's name"),
    (re.compile(r"(?i)(address|street|house_?no)"), "looks like an address"),
    (re.compile(r"(?i)(zip|postal|postcode|post_code|cep|pin_?code)"), "a postal code locates a person"),
    (re.compile(r"(?i)(^lat$|latitude|_lat$|^lng$|^lon$|longitude|_lng$|_lon$)"), "a coordinate locates a person"),
    (re.compile(r"(?i)(ip_?addr|^ip$|device_id|cookie)"), "an online identifier"),
    (re.compile(r"(?i)(birth|^dob$|_dob$)"), "a date of birth"),
    (re.compile(r"(?i)(ssn|passport|national_?id|aadhaar|^pan$|pan_number|tax_?id|licen[cs]e_?(no|number))"),
     "a government identifier"),
    (re.compile(r"(?i)(customer|client|user|patient|person|member|subscriber|employee|buyer|account_holder)"
                r"_?(unique_)?(id|key|number|no|uuid)$"), "identifies a person (pseudonymous ids are personal data)"),
]
SENSITIVE_RULES = [
    (re.compile(r"(?i)(comment|message|note|notes|feedback|free_?text|remark|body|review_text|complaint)"),
     "free text can contain anything a person typed"),
    (re.compile(r"(?i)(diagnos|condition|medication|disease|health|gender|race|ethnic|religio|sexual|"
                r"disabilit|salary|income)"), "special-category or confidential data"),
]


class AnalyzeError(ValueError):
    pass


def _qi(name: str) -> str:
    return '"' + str(name).replace('"', '""') + '"'


def _ident(text: str) -> str:
    s = re.sub(r"[^0-9a-zA-Z]+", "_", text).strip("_").lower()
    s = re.sub(r"_+", "_", s)
    if not s:
        s = "source"
    if s[0].isdigit():
        s = "s_" + s
    return s


# ---------------------------------------------------------------- discovery
def _batch_component(parts: list[str]) -> tuple[int, str, str] | None:
    """For a single file, find a path component that carries a batch id (batch_01, 2024-01, …)."""
    for i, comp in enumerate(parts):
        stem = comp.rsplit(".", 1)[0] if i == len(parts) - 1 else comp
        suffix = comp[len(stem):]
        m = DATE_TOKEN.search(stem)
        if m:
            return i, stem[:m.start()], stem[m.end():] + suffix
        m = re.search(r"\d+", stem)
        if m and BATCH_WORD.search(stem[:m.start()]):
            return i, stem[:m.start()], stem[m.end():] + suffix
    return None


def _trim_to_digits(prefix: str, suffix: str) -> tuple[str, str]:
    while prefix and prefix[-1].isdigit():
        prefix = prefix[:-1]
    while suffix and suffix[0].isdigit():
        suffix = suffix[1:]
    return prefix, suffix


def discover(project: Project, folder: str) -> dict[str, dict]:
    """Propose one source per distinct file shape. Files that differ only in a digit/date token
    (batch_01/x.csv, batch_02/x.csv or trips_2024-01.csv, trips_2024-02.csv) are one source
    whose location carries {batch}."""
    base = Path(folder) if Path(folder).is_absolute() else project.root / folder
    if not base.exists():
        raise AnalyzeError(f"folder not found: {folder}")
    files = sorted(p for p in base.rglob("*") if p.is_file() and p.suffix.lower() in DATA_EXT
                   and not any(part.startswith(".") for part in p.relative_to(base).parts))
    groups: dict[tuple, list[Path]] = {}
    for p in files:
        rel = p.relative_to(project.root) if p.is_relative_to(project.root) else p
        shape = tuple(re.sub(r"\d+", "#", c) for c in Path(rel).parts)
        groups.setdefault(shape, []).append(Path(rel))
    out: dict[str, dict] = {}
    for shape, members in groups.items():
        parts_list = [list(m.parts) for m in members]
        first = parts_list[0]
        location, batches = Path(*first).as_posix(), []
        if len(members) > 1:
            varying = [i for i in range(len(first)) if len({p[i] for p in parts_list}) > 1]
            if len(varying) == 1:
                i = varying[0]
                comps = [p[i] for p in parts_list]
                pre = _common_prefix(comps)
                suf = _common_prefix([c[::-1] for c in comps])[::-1]
                pre, suf = _trim_to_digits(pre, suf)
                batches = sorted(c[len(pre):len(c) - len(suf) if suf else None] for c in comps)
                new = list(first)
                new[i] = pre + "{batch}" + suf
                location = Path(*new).as_posix()
        else:
            hit = _batch_component(first)
            if hit:
                i, pre, suf = hit
                comp = first[i]
                tok = comp[len(pre):len(comp) - len(suf) if suf else None]
                new = list(first)
                new[i] = pre + "{batch}" + suf
                location, batches = Path(*new).as_posix(), [tok]
        stem = Path(location).name.rsplit(".", 1)[0].replace("{batch}", "")
        name = _ident(stem) or "source"
        while name in out:
            name += "_2"
        ext = Path(first[-1]).suffix.lower()
        out[name] = {"location": location, "format": DATA_EXT[ext], "batches_seen": batches,
                     "files": [m.as_posix() for m in members][:50]}
    return out


def _common_prefix(items: list[str]) -> str:
    if not items:
        return ""
    s1, s2 = min(items), max(items)
    for i, ch in enumerate(s1):
        if ch != s2[i]:
            return s1[:i]
    return s1


# ---------------------------------------------------------------- sample resolution
def sample_file(project: Project, spec: dict) -> Path | None:
    loc = str(spec.get("location") or "")
    if not loc or loc.lower().startswith(("http://", "https://")):
        return None
    pattern = loc.replace("{batch}", "*")
    base = project.root
    hits = sorted(base.glob(pattern)) if not Path(pattern).is_absolute() else sorted(Path("/").glob(pattern.lstrip("/")))
    hits = [h for h in hits if h.is_file()]
    return hits[0] if hits else None


# ---------------------------------------------------------------- profiling (counts only)
def _relation(con, path: Path, spec: dict, all_varchar: bool) -> str:
    fmt = str(spec.get("format", "")).lower()
    p = str(path).replace("'", "''")
    if fmt == "csv":
        c = spec.get("csv") if isinstance(spec.get("csv"), dict) else {}
        enc = str(c.get("encoding", "utf-8") or "utf-8").lower()
        if enc not in ("utf-8", "utf8"):
            from .bronze import _transcode
            p = str(_transcode(path, enc)).replace("'", "''")
        delim = str(c.get("delimiter", ",") or ",").replace("'", "''")
        quote = str(c.get("quote", '"') or '"').replace("'", "''")
        esc = str(c.get("escape", '"') or '"').replace("'", "''")
        header = "true" if C.as_bool(c.get("header", "yes") or "yes") else "false"
        nulls = [""] + [str(t) for t in C.as_list(c.get("null_tokens")) if not C.is_token(t, C.NA_TOKEN)]
        nullstr = "[" + ", ".join("'" + n.replace("'", "''") + "'" for n in nulls) + "]"
        av = "true" if all_varchar else "false"
        return (f"read_csv('{p}', delim='{delim}', quote='{quote}', escape='{esc}', header={header}, "
                f"all_varchar={av}, sample_size=20000, nullstr={nullstr}, ignore_errors=true)")
    if fmt == "parquet":
        return f"read_parquet('{p}')"
    if fmt in ("json", "jsonl"):
        return f"read_json_auto('{p}')"
    raise AnalyzeError(f"analysis supports csv, parquet and json, not {fmt}")


def _proposed_type(detected: str) -> str:
    t = str(detected).upper()
    base = t.split("(")[0]
    if base.startswith("DECIMAL"):
        return t
    if base in ("BIGINT", "INTEGER", "SMALLINT", "TINYINT", "HUGEINT", "UBIGINT", "UINTEGER"):
        return "BIGINT" if base in ("BIGINT", "HUGEINT", "UBIGINT") else "INTEGER"
    if base in ("DOUBLE", "FLOAT", "REAL"):
        return "DOUBLE"
    if base.startswith("TIMESTAMP"):
        return "TIMESTAMP"
    if base == "DATE":
        return "DATE"
    return "VARCHAR"  # BOOLEAN from Y/N, TIME, lists … stay text: decided in silver, never guessed


def profile_source(project: Project, name: str, spec: dict, path: Path) -> dict:
    import duckdb
    con = duckdb.connect()
    try:
        fmt = str(spec.get("format", "")).lower()
        if fmt == "xlsx":
            return _profile_xlsx(name, spec, path)
        raw = _relation(con, path, spec, all_varchar=True)
        con.execute(f"CREATE TEMP TABLE t AS SELECT * FROM {raw}")
        cols = [r[0] for r in con.execute("DESCRIBE t").fetchall()]
        rows = con.execute("SELECT COUNT(*) FROM t").fetchone()[0]
        distinct_rows = con.execute("SELECT COUNT(*) FROM (SELECT DISTINCT * FROM t)").fetchone()[0]
        try:
            typed = {r[0]: r[1] for r in con.execute(
                f"DESCRIBE SELECT * FROM {_relation(con, path, spec, all_varchar=False)}").fetchall()}
        except duckdb.Error:
            typed = {}
        prof_cols: dict[str, dict] = {}
        for c in cols:
            q = _qi(c)
            empty, distinct, lmin, lmax = con.execute(
                f"SELECT COUNT(*) FILTER (WHERE {q} IS NULL OR trim(CAST({q} AS VARCHAR)) = ''), "
                f"COUNT(DISTINCT {q}), MIN(length(CAST({q} AS VARCHAR))), MAX(length(CAST({q} AS VARCHAR))) "
                f"FROM t").fetchone()
            det = typed.get(c, "VARCHAR")
            info = {"detected_type": str(det), "empty": int(empty), "distinct": int(distinct),
                    "length_range": [lmin, lmax]}
            filled = rows - int(empty)
            if filled and (DATEISH_NAME.search(c) or str(det).upper().startswith(("TIMESTAMP", "DATE"))):
                fmts = []
                for fmtx in TS_FORMATS + DATE_FORMATS:
                    if fmtx == "%Y%m%d" and not DATEISH_NAME.search(c):
                        continue
                    n = con.execute(f"SELECT COUNT(*) FROM t WHERE try_strptime(trim(CAST({q} AS VARCHAR)), ?) "
                                    f"IS NOT NULL", [fmtx]).fetchone()[0]
                    if n:
                        fmts.append([fmtx, int(n)])
                fmts.sort(key=lambda x: -x[1])
                if fmts:
                    info["formats"] = fmts
                    info["unparsed"] = max(0, filled - sum(n for _, n in fmts))
            prof_cols[c] = info
        return {"sample": _rel(project, path), "rows": int(rows), "duplicate_rows": int(rows - distinct_rows),
                "columns": prof_cols}
    finally:
        con.close()


def _profile_xlsx(name: str, spec: dict, path: Path) -> dict:
    import openpyxl
    x = spec.get("xlsx") if isinstance(spec.get("xlsx"), dict) else {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[x.get("sheet")] if x.get("sheet") in wb.sheetnames else wb.worksheets[0]
    hdr = int(x.get("header_row", "1") or 1)
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = [str(h) if h is not None else f"column{i}" for i, h in enumerate(rows[hdr - 1])] if rows else []
    body = rows[hdr:]
    cols = {}
    for i, c in enumerate(header):
        vals = [r[i] if i < len(r) else None for r in body]
        filled = [v for v in vals if v is not None and str(v).strip() != ""]
        lens = [len(str(v)) for v in filled]
        cols[c] = {"detected_type": "VARCHAR", "empty": len(vals) - len(filled), "distinct": len(set(map(str, filled))),
                   "length_range": [min(lens) if lens else None, max(lens) if lens else None]}
    return {"sample": str(path.name), "rows": len(body), "duplicate_rows": len(body) - len(set(body)), "columns": cols}


def _rel(project: Project, path: Path) -> str:
    try:
        return path.relative_to(project.root).as_posix()
    except ValueError:
        return str(path)


def key_check(project: Project, spec: dict, path: Path, cols: list[str]) -> dict:
    """Is this column combination unique and never empty in the sample? (counts only)"""
    import duckdb
    con = duckdb.connect()
    try:
        con.execute(f"CREATE TEMP TABLE t AS SELECT * FROM {_relation(con, path, spec, all_varchar=True)}")
        tup = ", ".join(_qi(c) for c in cols)
        null_pred = " OR ".join(f"{_qi(c)} IS NULL OR trim({_qi(c)}) = ''" for c in cols)
        rows, distinct, nulls = con.execute(
            f"SELECT COUNT(*), (SELECT COUNT(*) FROM (SELECT DISTINCT {tup} FROM t)), "
            f"COUNT(*) FILTER (WHERE {null_pred}) FROM t").fetchone()
        return {"columns": cols, "duplicates": int(rows - distinct), "empty_keys": int(nulls)}
    finally:
        con.close()


def find_key(project: Project, spec: dict, path: Path, prof: dict) -> dict | None:
    rows = prof["rows"]
    cols = prof["columns"]
    if not rows:
        return None
    singles = [c for c, i in cols.items() if i["distinct"] == rows and i["empty"] == 0]
    id_like = [c for c in singles if ID_NAME.search(c)]
    if id_like:
        return {"columns": [id_like[0]], "duplicates": 0, "empty_keys": 0, "how": "unique id-like column"}
    cands = [c for c, i in cols.items() if (ID_NAME.search(c) or SEQ_NAME.search(c)) and i["empty"] == 0][:8]
    for i, a in enumerate(cands):
        for b in cands[i + 1:]:
            chk = key_check(project, spec, path, [a, b])
            if chk["duplicates"] == 0 and chk["empty_keys"] == 0:
                return {**chk, "how": "unique pair of id-like columns"}
    first = next(iter(cols), None)
    if first and first in singles:
        return {"columns": [first], "duplicates": 0, "empty_keys": 0, "how": "first column is unique"}
    return None


def run(project: Project, sources: dict[str, dict]) -> dict:
    """Profile every source that has a local sample. Returns and stores the analysis."""
    out = {"generated_at": now_iso(), "egress": "stats_only (counts, lengths, format matches; no values)",
           "sources": {}}
    for name, spec in sources.items():
        spec = spec if isinstance(spec, dict) else {}
        path = sample_file(project, spec)
        if path is None:
            out["sources"][name] = {"sample": None, "note": "no local sample file found for this location"}
            continue
        try:
            prof = profile_source(project, name, spec, path)
        except Exception as e:  # an unreadable sample must not stop the workbook
            out["sources"][name] = {"sample": _rel(project, path), "error": f"{type(e).__name__}: {str(e)[:200]}"}
            continue
        declared_key = [c for c, s in (C.get_path(spec, "schema.columns") or {}).items()
                        if isinstance(s, dict) and str(s.get("key", "")).lower() in C.TRUE_WORDS] \
            if isinstance(C.get_path(spec, "schema.columns"), dict) else []
        if declared_key and all(k in prof["columns"] for k in declared_key):
            prof["key"] = {**key_check(project, spec, path, declared_key), "how": "declared key"}
        else:
            found = find_key(project, spec, path, prof)
            if found:
                prof["key"] = found
        out["sources"][name] = prof
    C.atomic_write_json(project.loc("analysis"), out)
    audit(project, "intake.analyze", sources=len(sources))
    return out


def load(project: Project) -> dict:
    return C.read_json(project.loc("analysis"), {}) or {}


# ---------------------------------------------------------------- proposals from the analysis
def propose_type(col: str, info: dict) -> tuple[str, str]:
    det = _proposed_type(info.get("detected_type", "VARCHAR"))
    lr = info.get("length_range") or [None, None]
    ragged = lr[0] is not None and lr[1] is not None and lr[0] != lr[1]
    if det in ("BIGINT", "INTEGER", "DOUBLE") and (CODE_ALWAYS.search(col) or (CODE_IF_RAGGED.search(col) and ragged)):
        why = "a code, not a number — kept as text"
        if ragged:
            why += f" (values have {lr[0]}–{lr[1]} characters: leading zeros may have been lost)"
        return "VARCHAR", why
    if det == "DOUBLE" and MONEY_NAME.search(col):
        return "DECIMAL(18,2)", "money: exact decimal, never floating point"
    fm = info.get("formats") or []
    if fm and det == "VARCHAR":
        only_dates = all(f in DATE_FORMATS for f, _ in fm)
        return ("DATE" if only_dates else "TIMESTAMP"), "values parse as dates"
    if det == "TIMESTAMP" and fm and all(f in DATE_FORMATS for f, _ in fm):
        return "DATE", "values are dates without a time"
    return det, f"detected {det} in the sample"


def propose_classification(col: str, typ: str) -> tuple[str, str]:
    for rx, why in PII_RULES:
        if rx.search(col):
            return "pii", why
    if typ == "VARCHAR":
        for rx, why in SENSITIVE_RULES:
            if rx.search(col):
                return "sensitive", why
    else:
        for rx, why in SENSITIVE_RULES[1:]:
            if rx.search(col):
                return "sensitive", why
    return "internal", "no personal or special-category signal in the name"


def describe(info: dict, rows: int) -> str:
    """One stats-only line for the workbook's evidence column."""
    if not info:
        return ""
    e = info.get("empty", 0)
    pct = (100.0 * e / rows) if rows else 0
    s = f"{rows:,} rows · {e:,} empty ({pct:.1f}%) · {info.get('distinct', 0):,} distinct"
    lr = info.get("length_range") or [None, None]
    if lr[0] is not None and lr[0] != lr[1]:
        s += f" · length {lr[0]}–{lr[1]}"
    if info.get("formats"):
        s += " · formats: " + ", ".join(f"{f} ({n:,})" for f, n in info["formats"][:3])
        if info.get("unparsed"):
            s += f" · {info['unparsed']:,} unparsed"
    return s


def report(analysis: dict) -> str:
    lines = ["# Source analysis (counts only — no data values)", ""]
    for name, a in (analysis.get("sources") or {}).items():
        if a.get("error") or not a.get("rows") and a.get("sample") is None:
            lines.append(f"- **{name}**: {a.get('error') or a.get('note')}")
            continue
        key = a.get("key")
        ktxt = f"key {'+'.join(key['columns'])} ({key['how']}; {key['duplicates']} duplicate, " \
               f"{key['empty_keys']} empty)" if key else "no unique key found (the whole row can be the key)"
        lines.append(f"## {name} — {a.get('sample')}")
        lines.append(f"{a['rows']:,} rows · {a['duplicate_rows']:,} exact duplicate rows · {ktxt}")
        for c, info in a["columns"].items():
            t, _ = propose_type(c, info)
            lines.append(f"- `{c}` → {t}: {describe(info, a['rows'])}")
        lines.append("")
    return "\n".join(lines)
