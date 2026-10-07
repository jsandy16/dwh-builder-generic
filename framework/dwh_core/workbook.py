"""The intake workbook: every question of every layer in one Excel file, one tab per kind of
question, each answer pre-filled with a default the owner keeps or changes.

    dwh intake workbook [--out FILE]      write it (defaults from proposal.compose)
    dwh intake import FILE [--dry-run]    read it back: validate everything, then record

Import is all-or-nothing. The filled workbook is first applied to a scratch copy of the
project and every intake gate is run there; if anything is missing, invalid or
inconsistent, nothing is recorded and an "-issues" copy of the workbook is written with
every problem marked on its cell. Only a clean workbook is recorded, each answer
attributed to the person who answered for that role, with the cell it came from and
whether the default was kept.
"""
from __future__ import annotations

import copy
import hashlib
import json
import re
import shutil
import tempfile
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from . import VERSION, AGENT_IDS
from . import catalogue as K
from . import config as C
from . import proposal as PR
from .project import SPEC_FILES, Project, audit, now_iso

WB_VERSION = "1"
HDR, HINT, FIRST = 4, 5, 6  # header row, hint row, first data row

S_START, S_PEOPLE, S_PROJECT, S_GOV, S_POLICY = "Start here", "People", "Project", "Governance", "Reporting policy"
S_SOURCES, S_SCHEMA, S_CLASS, S_SILVER, S_ENT = "Sources", "Schema", "Classification", "Silver model", "Entities"
S_DQ, S_LOOK, S_KPI, S_GOLD, S_DASH, S_VIS = ("Data quality", "Lookups & flags", "KPIs", "Golden values",
                                              "Dashboard", "Dashboard visuals")
S_GATEB, S_META, S_ISSUES = "After first load", "_meta", "Issues"

ROLES = ["DE", "PO", "SME", "GOV"]
ROLE_NAMES = {"DE": "data engineer", "PO": "product owner", "SME": "business-rule owner", "GOV": "governance",
              "CON": "consumer owner"}
CHECK_NULL, CHECK_HARD, CHECK_ANOM, CHECK_STRUCT = "Empty values", "Impossible row", "Real but unusual", "Never filled"

# colours (Excel input-cell convention: blue text on yellow = something you may change)
NAVY, GREY_TXT, BLUE = "1F3864", "595959", "0000FF"
FILL_HDR, FILL_HINT, FILL_IN, FILL_REQ, FILL_INFO = "1F3864", "E7E6E6", "FFF2CC", "FCE4D6", "F2F2F2"
FILL_BAD = "FFC7CE"


class WorkbookError(ValueError):
    pass


# ================================================================ helpers
def _field(path_or_pattern: str) -> K.Field | None:
    f, _ = K.find_field(path_or_pattern)
    if f is None:
        for cat in K.all_catalogues():
            for g in cat.fields:
                if g.path == path_or_pattern:
                    return g
    return f


def _fmt(v: Any) -> str:
    """A spec value as cell text."""
    if v is None:
        return ""
    if isinstance(v, list):
        return ", ".join(_fmt(x) for x in v)
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False)
    return str(v)


def _cell(v: Any) -> str:
    """A cell value as answer text (Excel may hand back numbers, dates, booleans)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "yes" if v else "no"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d") if (v.hour, v.minute, v.second) == (0, 0, 0) else v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def _split(text: str) -> list[str]:
    if "\n" in text:
        parts = text.split("\n")
    else:
        parts = text.split(",")
    return [p.strip() for p in parts if p.strip()]


def _norm(label: Any) -> str:
    return re.sub(r"\s+", " ", str(label or "").replace("★", "")).strip().lower()


def _need(f: K.Field | None) -> str:
    if f is None:
        return "required"
    if f.level in ("M", "M*"):
        return "required"
    if f.level in ("C", "C*"):
        return "required when it applies"
    return "optional"


def _owner(f: K.Field | None, fallback: str = "DE") -> str:
    if f is None:
        return fallback
    return f.owner + (" ★" if f.human_owned else "")


def _allowed(f: K.Field | None) -> str:
    if f is None:
        return ""
    extra = []
    if f.level == "O":
        extra.append("NA")
    if f.none_only:
        extra.append("none")
    if f.human_owned and f.pending_ok:
        extra.append("pending")
    if f.choices:
        return ", ".join(list(f.choices) + extra)
    t = {"int": "whole number", "number": "number", "percent": "number 0–100", "bool": "yes / no",
         "date": "YYYY-MM-DD", "timezone": "IANA name, e.g. Asia/Kolkata", "duckdb_type":
         "VARCHAR, INTEGER, BIGINT, DOUBLE, DECIMAL(18,2), DATE, TIMESTAMP, BOOLEAN",
         "list": "comma-separated list", "columns": "comma-separated column names",
         "strftime_list": "comma-separated formats like %Y-%m-%d", "sql_predicate": "SQL condition",
         "measure": "count | sum:<col> | count_distinct:<col>", "person": "a person id from People",
         "column": "a column name", "text": "text", "path_pattern": "path with {batch}"}.get(f.type, "text")
    return t + (f" (or {', '.join(extra)})" if extra else "")


def _dv_choices(f: K.Field | None) -> list[str] | None:
    if f is None or not f.choices or f.type in ("list", "columns"):
        return None
    out = list(f.choices)
    if f.level == "O":
        out.append("NA")
    if f.human_owned and f.pending_ok:
        out.append("pending")
    return out


# ================================================================ writing
@dataclass
class Col:
    key: str
    label: str
    pattern: str = ""            # catalogue pattern for owner/level/choices
    kind: str = "answer"         # answer | info | key
    width: int = 16
    hint: str = ""
    choices: list | None = None
    wrap: bool = False
    required: bool | None = None  # None: from the catalogue field (no field = not required)


class Writer:
    def __init__(self, project: Project, prop: PR.Proposal):
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        self.openpyxl = openpyxl
        self.Font, self.Fill, self.Align = Font, PatternFill, Alignment
        thin = Side(style="thin", color="BFBFBF")
        self.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        self.project, self.prop = project, prop
        self.wb = openpyxl.Workbook()
        self.wb.remove(self.wb.active)
        self.defaults: dict[str, Any] = {}
        self.counts: dict[str, list[int]] = {}   # tab -> [questions, required blanks]
        recorded = {ns: project.load_namespace(ns) for ns in SPEC_FILES if ns not in ("synthetic",)}
        self.answers = overlay(prop.doc, {k: v for k, v in recorded.items() if v})
        if isinstance(recorded.get("project"), dict):
            for k in ("dwh_core_version", "key_algo", "layout"):
                self.answers.get("project", {}).pop(k, None)

    # ---- styles
    def font(self, **kw):
        base = {"name": "Arial", "size": 10}
        base.update(kw)
        return self.Font(**base)

    def fill(self, rgb):
        return self.Fill("solid", start_color=rgb, end_color=rgb)

    def sheet(self, title: str, heading: str, intro: str):
        ws = self.wb.create_sheet(title)
        ws.sheet_view.showGridLines = False
        ws["A1"] = heading
        ws["A1"].font = self.font(size=14, bold=True, color=NAVY)
        ws["A2"] = intro
        ws["A2"].font = self.font(color=GREY_TXT)
        ws["A2"].alignment = self.Align(wrap_text=True, vertical="top")
        ws.row_dimensions[2].height = 42
        self.counts.setdefault(title, [0, 0])
        return ws

    def header(self, ws, cols: list[Col], intro_span: int | None = None):
        from openpyxl.utils import get_column_letter
        span = intro_span or len(cols)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(1, min(span, 12)))
        for i, c in enumerate(cols, start=1):
            f = _field(c.pattern) if c.pattern else None
            label = c.label + (" ★" if f is not None and f.human_owned else "")
            h = ws.cell(row=HDR, column=i, value=label)
            h.font = self.font(bold=True, color="FFFFFF")
            h.fill = self.fill(FILL_HDR)
            h.alignment = self.Align(wrap_text=True, vertical="center")
            h.border = self.border
            if c.kind == "info":
                hint = c.hint or "information only"
            elif c.kind == "key":
                hint = c.hint or "name (identifier)"
            else:
                hint = c.hint or " · ".join(x for x in (_owner(f), _need(f), _allowed(f)) if x)
            t = ws.cell(row=HINT, column=i, value=hint)
            t.font = self.font(italic=True, size=8, color=GREY_TXT)
            t.fill = self.fill(FILL_HINT)
            t.alignment = self.Align(wrap_text=True, vertical="top")
            t.border = self.border
            ws.column_dimensions[get_column_letter(i)].width = c.width
        ws.row_dimensions[HDR].height = 30
        ws.row_dimensions[HINT].height = 48

    def put(self, ws, row: int, col: int, value: Any, kind: str = "answer", required: bool = False,
            wrap: bool = False, example: bool = False):
        cell = ws.cell(row=row, column=col, value=_fmt(value) if value is not None else None)
        cell.border = self.border
        cell.alignment = self.Align(wrap_text=wrap, vertical="top")
        if example:
            cell.font = self.font(italic=True, color="808080")
            return cell
        if kind == "answer":
            cell.number_format = "@"
            cell.font = self.font(color=BLUE)
            blank = value is None or _fmt(value) == ""
            cell.fill = self.fill(FILL_REQ if (required and blank) else FILL_IN)
        elif kind == "info":
            cell.font = self.font(color=GREY_TXT, size=9)
            cell.fill = self.fill(FILL_INFO)
        else:
            cell.font = self.font(bold=True)
        return cell

    def validate(self, ws, cells: dict[tuple, list[str]], strict: bool = True):
        """cells: (choices tuple) -> list of coordinates."""
        from openpyxl.worksheet.datavalidation import DataValidation
        for choices, coords in cells.items():
            if not coords:
                continue
            formula = '"' + ",".join(choices) + '"'
            if len(formula) > 255:
                continue
            dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                                errorStyle="stop" if strict else "warning", showErrorMessage=True,
                                error="Pick a value from the list (or check the hint row).",
                                errorTitle="Not an allowed answer")
            ws.add_data_validation(dv)
            for co in coords:
                dv.add(co)

    def remember(self, path: str, value: Any):
        if value is not None and path not in self.defaults:
            self.defaults[path] = value

    def ans(self, path: str) -> Any:
        return C.get_path(self.answers, path)

    def dflt(self, path: str) -> Any:
        v = C.get_path(self.prop.doc, path)
        self.remember(path, v)
        return v

    # ---- Q&A tab
    def qa(self, title: str, heading: str, intro: str, items: list[tuple]):
        """items: (path, question override or '', owner override or '', choices override or None)."""
        ws = self.sheet(title, heading, intro)
        cols = [Col("id", "Question ID", kind="info", width=26, hint="do not change"),
                Col("q", "Question", kind="info", width=58, hint="what is asked"),
                Col("owner", "Owner", kind="info", width=10, hint="role · ★ = owner decision"),
                Col("need", "Need", kind="info", width=13, hint="required / optional"),
                Col("allowed", "Allowed answers", kind="info", width=30, hint="pick or type"),
                Col("default", "Default", kind="info", width=20, hint="the proposal"),
                Col("answer", "Your answer", width=26, hint="keep the default or change it"),
                Col("why", "Why it matters / why this default", kind="info", width=60, hint="")]
        self.header(ws, cols)
        dv: dict[tuple, list[str]] = {}
        loose: dict[tuple, list[str]] = {}
        r = FIRST
        for path, q, owner, choices in items:
            pseudo = path.startswith("answered_by.")
            f = None if pseudo else _field(path)
            if pseudo:
                question, default = q
                answer = default
                self.remember(path, default)
            else:
                question = q or (f.prompt if f else "")
                default = self.dflt(path)
                answer = self.ans(path)
                if answer is None:
                    answer = default
                if f is not None and f.path.split(".")[-1] in PR.NO_DEFAULT:
                    default = answer = None
            required = True if pseudo else _need(f) == "required"
            allowed = "a person id from the People tab" if pseudo else _allowed(f)
            vals = [path, question, owner or ("DE" if pseudo else _owner(f)), "required" if pseudo else _need(f),
                    allowed, default, answer, ("recorded as the author of that role's answers" if pseudo else
                                               (self.prop.note(path) or (f.why if f else "")))]
            for i, v in enumerate(vals, start=1):
                self.put(ws, r, i, v, kind=cols[i - 1].kind, required=required, wrap=i in (2, 5, 8))
            ch = choices if pseudo else (choices or _dv_choices(f))
            if ch:
                (loose if pseudo else dv).setdefault(tuple(ch), []).append(f"G{r}")
            self.counts[title][0] += 1
            if required and _fmt(answer) == "":
                self.counts[title][1] += 1
            ws.row_dimensions[r].height = 30
            r += 1
        self.validate(ws, loose, strict=False)
        self.validate(ws, dv)
        ws.freeze_panes = f"C{FIRST}"
        return ws

    # ---- grid tab
    def grid(self, title: str, heading: str, intro: str, cols: list[Col], rows: list[dict],
             example: dict | None = None, row_choices: dict | None = None, freeze: str = "C"):
        ws = self.sheet(title, heading, intro)
        self.header(ws, cols)
        dv: dict[tuple, list[str]] = {}
        loose: dict[tuple, list[str]] = {}
        r = FIRST
        if example:
            for i, c in enumerate(cols, start=1):
                self.put(ws, r, i, example.get(c.key, "example" if i == 1 else None), example=True)
            r += 1
        from openpyxl.utils import get_column_letter
        for row in rows:
            for i, c in enumerate(cols, start=1):
                v = row.get(c.key)
                pat = row.get("_pattern", {}).get(c.key) or c.pattern
                f = _field(pat) if pat else None
                base_req = c.required if c.required is not None else (f is not None and _need(f) == "required")
                required = c.kind == "answer" and row.get("_required", {}).get(c.key, base_req)
                self.put(ws, r, i, v, kind=c.kind, required=bool(required), wrap=c.wrap)
                coord = f"{get_column_letter(i)}{r}"
                ch = (row.get("_choices") or {}).get(c.key) or c.choices or (_dv_choices(f) if c.kind == "answer" else None)
                if ch and c.kind == "answer":
                    target = loose if (row.get("_loose") or {}).get(c.key) else dv
                    target.setdefault(tuple(ch), []).append(coord)
                if c.kind == "answer":
                    self.counts[title][0] += 1
                    if required and _fmt(v) == "":
                        self.counts[title][1] += 1
            r += 1
        # spare input rows for additions
        for _ in range(row_choices.get("_spare", 0) if row_choices else 0):
            for i, c in enumerate(cols, start=1):
                self.put(ws, r, i, None, kind="answer" if c.kind != "info" else "info")
                if c.choices and c.kind == "answer":
                    dv.setdefault(tuple(c.choices), []).append(f"{get_column_letter(i)}{r}")
            r += 1
        self.validate(ws, dv)
        self.validate(ws, loose, strict=False)
        ws.freeze_panes = f"{freeze}{FIRST}"
        ws.auto_filter.ref = f"A{HDR}:{get_column_letter(len(cols))}{max(r - 1, HDR)}"
        return ws


# Maps whose keys are chosen by people (sources, rules, KPIs …). Once answered, the recorded keys
# are the whole collection: an item the owner removed is not brought back by the proposal.
COLLECTIONS = ["people", "sources", "sources.*.schema.columns", "silver.entities", "silver.entities.*.columns",
               "silver.entities.*.mapping", "silver.entities.*.mapping.*", "silver.entities.*.null_policy",
               "silver.entities.*.structural_nulls", "silver.entities.*.valid_anomalies",
               "silver.entities.*.hard_rejects", "silver.entities.*.lookups", "silver.entities.*.flags",
               "silver.entities.*.masking", "silver.entities.*.drop_columns", "metrics", "serve.kpis", "serve.charts"]


def _is_collection(path: str) -> bool:
    parts = C.split_path(path)
    for pat in COLLECTIONS:
        pp = C.split_path(pat)
        if len(pp) == len(parts) and all(a == "*" or a == b for a, b in zip(pp, parts)):
            return True
    return False


def overlay(prop: Any, rec: Any, path: str = "") -> Any:
    """Recorded answers over the proposal: leaves the owner never answered keep their default,
    but for collections the recorded keys win (no resurrection of removed items)."""
    if rec is None:
        return copy.deepcopy(prop)
    if not isinstance(rec, dict) or not isinstance(prop, dict):
        return copy.deepcopy(rec)
    keys = list(rec) if _is_collection(path) else list(dict.fromkeys(list(prop) + list(rec)))
    return {k: overlay(prop.get(k), rec.get(k), f"{path}.{k}" if path else k) for k in keys}


def _role_holders(people: dict) -> dict[str, list[str]]:
    out: dict[str, list[str]] = {}
    for pid, rec in (people or {}).items() if isinstance(people, dict) else []:
        for r in C.as_list((rec or {}).get("roles") if isinstance(rec, dict) else []):
            out.setdefault(str(r).strip().upper(), []).append(pid)
    return out


def build(project: Project, out: Path | None = None, run_analysis: bool = True) -> tuple[Path, dict]:
    prop = PR.compose(project, run_analysis=run_analysis)
    w = Writer(project, prop)
    A = w.answers
    sources = A.get("sources") if isinstance(A.get("sources"), dict) else {}
    ents = C.get_path(A, "silver.entities") or {}
    metrics = A.get("metrics") if isinstance(A.get("metrics"), dict) else {}
    analysis = prop.analysis.get("sources") or {}

    _start_sheet(w)  # filled in at the end, once counts are known

    # ---------------- People
    ppl = A.get("people") if isinstance(A.get("people"), dict) else {}
    rows = [{"id": pid, "name": (rec or {}).get("name"), "email": (rec or {}).get("email") or "NA",
             "roles": C.as_list((rec or {}).get("roles"))} for pid, rec in ppl.items()]
    for pid, rec in ppl.items():
        for k in ("name", "email", "roles"):
            w.remember(f"people.{pid}.{k}", C.get_path(prop.doc, f"people.{pid}.{k}"))
    w.grid(S_PEOPLE, "People — who answers and approves",
           "One row per person. Roles (comma-separated): DE = data engineer, PO = product owner, SME = business-rule "
           "owner, GOV = governance, CON = consumer owner. One person may hold several roles. Only people listed here "
           "can give the ★ answers of their role. The assistant is never a person here.",
           [Col("id", "Person id", kind="answer", width=16, required=True, hint="DE · required · short id, e.g. sandeep"),
            Col("name", "Full name", "people.*.name", width=26),
            Col("email", "Email", "people.*.email", width=30),
            Col("roles", "Roles", "people.*.roles", width=24, hint="DE · required · comma-separated: DE, PO, SME, GOV, CON")],
           rows, example={"id": "example", "name": "Dana Rao", "email": "dana@example.com", "roles": "DE, SME"},
           row_choices={"_spare": 4})

    # ---------------- Project (+ who answered each role)
    holders = _role_holders(ppl)
    items = [("project.name", "", "", None), ("project.profile", "", "", None), ("project.builder", "", "", None)]
    for role in ROLES:
        who = (holders.get(role) or [""])[0]
        items.append((f"answered_by.{role}",
                      (f"Who answered the {role} ({ROLE_NAMES[role]}) questions in this workbook? "
                       f"Their name is recorded on every {role} answer.", who),
                      "DE", sorted(ppl) or None))
    w.qa(S_PROJECT, "Project", "The project's name and profile, and who answered each role's questions in this "
         "workbook (answers are recorded in that person's name).", items)

    # ---------------- Governance
    w.qa(S_GOV, "Governance ★", "Decisions for the governance owner (GOV). ★ = an owner decision: the default is a "
         "proposal — keeping it records it as the owner's own answer.",
         [("policies.compliance", "", "", None), ("policies.egress", "", "", None),
          ("serve.row_security", "", "", ["none", "pending"])])

    # ---------------- Reporting policy
    w.qa(S_POLICY, "Reporting policy ★", "Decisions for the product owner (PO): which day a timestamp belongs to, "
         "the calendar and money precision. Every KPI depends on these.",
         [("policies.timezone.reporting", "", "", None), ("policies.calendar.type", "", "", None),
          ("policies.calendar.week_start", "", "", None), ("policies.money.scale", "", "", None)])

    # ---------------- Sources
    scols = [Col("source", "Source", kind="key", width=20, hint="source name (identifier) — renaming means renaming it on every tab"),
             Col("sample", "Sample analysed", kind="info", width=30), Col("rows_seen", "Rows in sample", kind="info", width=11),
             Col("role", "Role", "sources.*.role", width=12), Col("connector", "Connector", "sources.*.connector", width=12),
             Col("location", "Location", "sources.*.location", width=44, wrap=True),
             Col("format", "Format", "sources.*.format", width=9),
             Col("load_type", "Load type", "sources.*.load_type", width=18),
             Col("redelivery_policy", "Re-delivered batch", "sources.*.redelivery_policy", width=13),
             Col("drift_policy", "Unexpected new columns", "sources.*.drift_policy", width=14),
             Col("freshness.basis", "Freshness basis", "sources.*.freshness.basis", width=13),
             Col("freshness.token_format", "Batch id format", "sources.*.freshness.token_format", width=12),
             Col("freshness.column", "Freshness column", "sources.*.freshness.column", width=16),
             Col("freshness.formats", "Freshness column formats", "sources.*.freshness.formats", width=18),
             Col("freshness.window_days", "Max age (days)", "sources.*.freshness.window_days", width=10),
             Col("volume.min_rows", "Min rows", "sources.*.volume.min_rows", width=9),
             Col("volume.max_rows", "Max rows", "sources.*.volume.max_rows", width=9),
             Col("volume.threshold_source", "Row bounds from", "sources.*.volume.threshold_source", width=12),
             Col("retry_attempts", "Retries", "sources.*.retry_attempts", width=8),
             Col("credentials_env", "Credentials env var", "sources.*.credentials_env", width=14),
             Col("profile_split_by", "Profile split by", "sources.*.profile_split_by", width=14),
             Col("csv.delimiter", "CSV delimiter", "sources.*.csv.delimiter", width=9),
             Col("csv.header", "CSV header row", "sources.*.csv.header", width=9),
             Col("csv.quote", "CSV quote", "sources.*.csv.quote", width=8),
             Col("csv.escape", "CSV escape", "sources.*.csv.escape", width=8),
             Col("csv.encoding", "CSV encoding", "sources.*.csv.encoding", width=10),
             Col("csv.null_tokens", "CSV null tokens", "sources.*.csv.null_tokens", width=12),
             Col("xlsx.sheet", "Excel sheet", "sources.*.xlsx.sheet", width=12),
             Col("xlsx.header_row", "Excel header row", "sources.*.xlsx.header_row", width=10)]
    rows = []
    for s in sources:
        an = analysis.get(s) or {}
        row = {"source": s, "sample": an.get("sample") or an.get("note") or an.get("error") or "",
               "rows_seen": f"{an['rows']:,}" if an.get("rows") is not None else "", "_required": {}}
        fmt = str(C.get_path(A, f"sources.{s}.format") or "").lower()
        for c in scols[3:]:
            p = f"sources.{s}.{c.key}"
            row[c.key] = w.ans(p)
            w.dflt(p)
            f = _field(c.pattern)
            if f is not None and f.level in ("C", "C*"):
                row["_required"][c.key] = K.eval_condition(f.condition, A, [s])
        if fmt != "csv":
            for c in scols:
                if c.key.startswith("csv."):
                    row["_required"][c.key] = False
        rows.append(row)
    w.grid(S_SOURCES, "Sources — landing the files (bronze)",
           "One row per source feed. Batches are found by putting {batch} where the batch id appears in the location. "
           "Columns that do not apply (e.g. CSV settings for a Parquet source, freshness details when the basis is none) "
           "can stay blank.", scols, rows, freeze="B")

    # ---------------- Schema
    sc = [Col("source", "Source", kind="key", width=18), Col("column", "Column", kind="key", width=30),
          Col("detected", "Detected in sample", kind="info", width=13),
          Col("evidence", "Evidence (counts only)", kind="info", width=44, wrap=True),
          Col("type", "Type", "sources.*.schema.columns.*.type", width=16),
          Col("required", "Required in every file", "sources.*.schema.columns.*.required", width=11, choices=["yes", "no"]),
          Col("key", "Part of the record key", "sources.*.schema.columns.*.key", width=11, choices=["yes", "no", "NA"]),
          Col("allowed_values", "Allowed values", "sources.*.schema.columns.*.allowed_values", width=24),
          Col("why", "Why this default", kind="info", width=40, wrap=True)]
    rows = []
    for s in sources:
        an = analysis.get(s) or {}
        for c in (C.get_path(A, f"sources.{s}.schema.columns") or {}):
            info = (an.get("columns") or {}).get(c) or {}
            base = f"sources.{s}.schema.columns.{c}"
            row = {"source": s, "column": c, "detected": info.get("detected_type", ""),
                   "evidence": _describe(info, an.get("rows"))}
            for k in ("type", "required", "key", "allowed_values"):
                row[k] = w.ans(f"{base}.{k}")
                w.dflt(f"{base}.{k}")
            row["why"] = prop.note(f"{base}.type") or prop.note(f"{base}.key")
            rows.append(row)
    w.grid(S_SCHEMA, "Schema — the contract every file must meet (mandatory)",
           "One row per column of each source. The schema can never be NA. Types were detected from the sample; money is "
           "DECIMAL, codes with leading zeros are text. Add a row for a column that will appear in later files.",
           sc, rows, freeze="C")

    # ---------------- Classification (+ masking in silver)
    cc = [Col("source", "Source", kind="key", width=18), Col("column", "Column", kind="key", width=30),
          Col("classification", "Classification", "sources.*.schema.columns.*.classification", width=14),
          Col("masking", "Handling in silver", "silver.entities.*.masking.*", width=13,
              hint="GOV ★ · required for pii / sensitive / regulated · hash (joinable, unreadable) / drop / keep"),
          Col("why", "Why this default", kind="info", width=50, wrap=True)]
    rows = []
    src_entity = _source_entity_map(A)
    for s in sources:
        for c in (C.get_path(A, f"sources.{s}.schema.columns") or {}):
            base = f"sources.{s}.schema.columns.{c}"
            e, cn = src_entity.get((s, c), (None, None))
            mpath = f"silver.entities.{e}.masking.{cn}" if e else ""
            cls = w.ans(f"{base}.classification")
            w.dflt(f"{base}.classification")
            mask = w.ans(mpath) if mpath else None
            if mpath:
                w.dflt(mpath)
            rows.append({"source": s, "column": c, "classification": cls, "masking": mask,
                         "why": prop.note(f"{base}.classification"),
                         "_required": {"masking": str(cls).lower() in ("pii", "sensitive", "regulated")}})
    w.grid(S_CLASS, "Classification ★ — governance decides how each column is protected",
           "One row per column. public / internal need no handling; pii / sensitive / regulated must be hashed, dropped or "
           "explicitly kept when the data moves to silver. The defaults come from the column names only — check them.",
           cc, rows, freeze="C")

    # ---------------- Silver model
    sm = [Col("source", "Source", kind="key", width=16), Col("column", "Source column", kind="key", width=28),
          Col("entity", "Entity", width=16, required=True, hint="DE · required · silver entity this column goes to"),
          Col("silver_column", "Silver column", width=28, required=True, hint="DE · required · canonical name (letters, digits, _)"),
          Col("expression", "Expression (optional)", width=34, wrap=True,
              hint="DE · optional · SQL over source columns; blank = copy the column"),
          Col("type", "Silver type", "silver.entities.*.columns.*.type", width=15),
          Col("formats", "Date/time formats", "silver.entities.*.columns.*.formats", width=22),
          Col("source_timezone", "Recorded in time zone", "silver.entities.*.columns.*.source_timezone", width=18),
          Col("drop", "Not carried — reason", width=22, hint="DE · optional · fill to leave this column out of silver")]
    rows = []
    for e, spec in ents.items():
        mapping = spec.get("mapping") if isinstance(spec.get("mapping"), dict) else {}
        drops = spec.get("drop_columns") if isinstance(spec.get("drop_columns"), dict) else {}
        for s in C.as_list(spec.get("sources")):
            m = mapping.get(s) if isinstance(mapping.get(s), dict) else {}
            rev = {str(v): k for k, v in m.items()}
            for c in (C.get_path(A, f"sources.{s}.schema.columns") or {}):
                if c in drops:
                    rows.append({"source": s, "column": c, "entity": e, "drop": drops[c]})
                    continue
                cn = rev.get(c, PR.canon(c))
                expr = m.get(cn)
                expr = "" if expr in (None, c) else expr
                base = f"silver.entities.{e}.columns.{cn}"
                row = {"source": s, "column": c, "entity": e, "silver_column": cn, "expression": expr,
                       "type": w.ans(f"{base}.type"), "formats": w.ans(f"{base}.formats"),
                       "source_timezone": w.ans(f"{base}.source_timezone"), "_required": {}}
                for k in ("type", "formats", "source_timezone"):
                    w.dflt(f"{base}.{k}")
                t = str(row["type"] or "").upper()
                row["_required"] = {"formats": t in ("DATE", "TIMESTAMP"), "source_timezone": t == "TIMESTAMP"}
                rows.append(row)
            for cn, expr in m.items():  # derived columns not tied to one source column
                if str(expr).strip() not in (C.get_path(A, f"sources.{s}.schema.columns") or {}) and \
                        str(expr).strip().upper() != "NULL":
                    if any(r.get("silver_column") == cn and r.get("source") == s for r in rows):
                        continue
                    base = f"silver.entities.{e}.columns.{cn}"
                    rows.append({"source": s, "column": "", "entity": e, "silver_column": cn, "expression": expr,
                                 "type": w.ans(f"{base}.type"), "formats": w.ans(f"{base}.formats"),
                                 "source_timezone": w.ans(f"{base}.source_timezone")})
    w.grid(S_SILVER, "Silver model — where each column goes, and its silver type",
           "One row per source column: the entity and canonical column it feeds, an optional SQL expression (e.g. "
           "lpad(zip, 5, '0')), the silver type, every date format present, and the time zone timestamps were recorded "
           "in. Several sources may feed one entity; give the same silver column name to conform them. Add a row with "
           "no source column for a column computed from others.",
           sm, rows, freeze="C")

    # ---------------- Entities
    ec = [Col("entity", "Entity", kind="key", width=16), Col("fed_by", "Fed by", kind="info", width=18),
          Col("dupes", "Key check in sample", kind="info", width=26, wrap=True),
          Col("natural_key", "Record key", "silver.entities.*.natural_key", width=26),
          Col("tiebreaker", "Tiebreaker", "silver.entities.*.tiebreaker", width=26),
          Col("merge.strategy", "Merge strategy", "silver.entities.*.merge.strategy", width=17),
          Col("merge.version_column", "Version column", "silver.entities.*.merge.version_column", width=15),
          Col("merge.partition_column", "Partition column", "silver.entities.*.merge.partition_column", width=12),
          Col("merge.op_column", "CDC op column", "silver.entities.*.merge.op_column", width=11),
          Col("merge.op_codes.insert", "CDC insert code", "silver.entities.*.merge.op_codes.insert", width=9),
          Col("merge.op_codes.update", "CDC update code", "silver.entities.*.merge.op_codes.update", width=9),
          Col("merge.op_codes.delete", "CDC delete code", "silver.entities.*.merge.op_codes.delete", width=9),
          Col("merge.sequence_column", "CDC sequence column", "silver.entities.*.merge.sequence_column", width=12),
          Col("rejected_survivor_policy", "If the newest version breaks a rule",
              "silver.entities.*.rejected_survivor_policy", width=16),
          Col("dead_letter_tolerance_pct", "Max % of a batch rejected", "silver.entities.*.dead_letter_tolerance_pct", width=11),
          Col("tolerance_source", "Tolerance from", "silver.entities.*.tolerance_source", width=11),
          Col("target_table", "Silver table name", "silver.entities.*.target_table", width=16),
          Col("why", "Why these defaults", kind="info", width=44, wrap=True)]
    rows = []
    for e, spec in ents.items():
        srcs = C.as_list(spec.get("sources"))
        key_txt = []
        for s in srcs:
            k = (analysis.get(s) or {}).get("key")
            if k:
                key_txt.append(f"{s}: {'+'.join(k['columns'])} — {k['duplicates']:,} duplicate, {k['empty_keys']:,} empty")
            elif (analysis.get(s) or {}).get("rows") is not None:
                key_txt.append(f"{s}: no unique key; {(analysis.get(s) or {}).get('duplicate_rows', 0):,} exact duplicate rows")
        row = {"entity": e, "fed_by": ", ".join(srcs), "dupes": "\n".join(key_txt), "_required": {}}
        for c in ec[3:-1]:
            p = f"silver.entities.{e}.{c.key}"
            row[c.key] = w.ans(p)
            w.dflt(p)
            f = _field(c.pattern)
            if f is not None and f.level in ("C", "C*"):
                row["_required"][c.key] = K.eval_condition(f.condition, A, [e])
        row["why"] = "; ".join(x for x in (prop.note(f"silver.entities.{e}.natural_key"),
                                           prop.note(f"silver.entities.{e}.merge.strategy")) if x)
        rows.append(row)
    w.grid(S_ENT, "Entities — identity, duplicates and how batches merge",
           "One row per silver entity. The record key identifies one record (identity columns only, never a measure). "
           "The tiebreaker decides which duplicate wins. CDC columns only apply to the cdc_apply strategy; the version "
           "column only to upsert_by_version.", ec, rows, freeze="B")

    # ---------------- Data quality
    dq = [Col("entity", "Entity", kind="key", width=15),
          Col("check", "Check", kind="key", width=16, hint="kind of check"),
          Col("name", "Column / rule name", width=30, required=False, hint="column, or a short rule name (letters, digits, _)"),
          Col("condition", "Condition (SQL)", width=46, wrap=True, required=False, hint="SME · rows the rule matches, on silver columns"),
          Col("reason", "Reason", width=30, wrap=True, required=False,
              hint="SME · impossible row: short code (e.g. dropoff_before_pickup) · real but unusual: why it is real"),
          Col("evidence", "Evidence (counts only)", kind="info", width=40, wrap=True),
          Col("decision", "Decision", width=16, required=True, hint="see the choices in the cell's list")]
    rows = []
    for e, spec in ents.items():
        base = f"silver.entities.{e}"
        for kind, key, star in ((CHECK_HARD, "hard_rejects", False), (CHECK_ANOM, "valid_anomalies", True)):
            rules = w.ans(f"{base}.{key}")
            w.dflt(f"{base}.{key}")
            if isinstance(rules, dict) and rules:
                for rn, rr in rules.items():
                    rr = rr if isinstance(rr, dict) else {}
                    rows.append({"entity": e, "check": kind, "name": rn, "condition": rr.get("predicate"),
                                 "reason": rr.get("reason"), "evidence": prop.note(f"{base}.{key}.{rn}"),
                                 "decision": rr.get("decision", "apply"), "_choices": {"decision": ["apply", "skip"]},
                                 "_pattern": {"decision": f"silver.entities.*.{key}"}})
            else:
                rows.append({"entity": e, "check": kind, "name": "", "condition": "", "reason": "",
                             "evidence": ("Rows that are physically impossible go to the dead-letter file with this reason."
                                          if key == "hard_rejects" else
                                          "Rows that look wrong but are real business events are always kept."),
                             "decision": "none", "_choices": {"decision": ["none", "apply", "skip"]},
                             "_pattern": {"decision": f"silver.entities.*.{key}"},
                             "_required": {"name": False, "condition": False, "reason": False}})
        sn = w.ans(f"{base}.structural_nulls")
        if isinstance(sn, dict):
            for cand, dec in sn.items():
                p = f"{base}.structural_nulls.{cand}"
                w.dflt(p)
                rows.append({"entity": e, "check": CHECK_STRUCT, "name": cand, "condition": "",
                             "reason": "", "evidence": prop.note(p), "decision": dec,
                             "_choices": {"decision": ["structural", "defect"]},
                             "_pattern": {"decision": "silver.entities.*.structural_nulls.*"},
                             "_required": {"condition": False, "reason": False}})
        np_ = w.ans(f"{base}.null_policy") or {}
        for c, pol in (np_.items() if isinstance(np_, dict) else []):
            p = f"{base}.null_policy.{c}"
            w.dflt(p)
            rows.append({"entity": e, "check": CHECK_NULL, "name": c, "condition": "", "reason": "",
                         "evidence": _null_evidence(A, analysis, e, c) or prop.note(p), "decision": pol,
                         "_choices": {"decision": ["drop", "keep", "impute:0"]}, "_loose": {"decision": True},
                         "_pattern": {"decision": "silver.entities.*.null_policy.*"},
                         "_required": {"condition": False, "reason": False}})
    w.grid(S_DQ, "Data quality — rules, empty values and never-filled columns",
           "Impossible row: dead-lettered with a reason (SME). Real but unusual ★: rows that look wrong but must be kept "
           "(e.g. refunds). For each entity, keep 'none' if there are none, or replace that row with a rule (name, SQL "
           "condition, reason) and set Decision to apply; add rows for more rules. Empty values: drop (dead-letter with a "
           "reason) / keep (stays empty) / impute:<value> (filled and flagged). Never filled ★: a column that is never "
           "filled — structural (not measured, never imputed) or defect (missing; the empty-value policy applies).",
           dq, rows, example={"entity": "example", "check": CHECK_HARD, "name": "dropoff_before_pickup",
                              "condition": "dropoff_ts < pickup_ts", "reason": "impossible_time_order",
                              "evidence": "rows like this are dead-lettered", "decision": "apply"},
           freeze="C")

    # ---------------- Lookups & flags
    lc = [Col("entity", "Entity", width=15, required=False, hint="SME/DE · the entity being enriched"),
          Col("kind", "Kind", width=9, choices=["lookup", "flag"], required=False, hint="lookup / flag"),
          Col("name", "Name", width=18, required=False, hint="short name (letters, digits, _)"),
          Col("reference", "Reference entity", "silver.entities.*.lookups.*.reference", width=16),
          Col("on", "Join on", "silver.entities.*.lookups.*.on", width=30,
              hint="DE · lookup · local_column = reference_column (comma-separated pairs)"),
          Col("columns", "Columns to bring", "silver.entities.*.lookups.*.columns", width=30),
          Col("missing", "If the code is missing", "silver.entities.*.lookups.*.missing", width=13),
          Col("condition", "Flag condition (SQL)", "silver.entities.*.flags.*", width=40, wrap=True),
          Col("why", "Why", kind="info", width=40, wrap=True)]
    rows = []
    for e, spec in ents.items():
        base = f"silver.entities.{e}"
        lk = w.ans(f"{base}.lookups")
        w.dflt(f"{base}.lookups")
        for n, l in (lk.items() if isinstance(lk, dict) else []):
            l = l if isinstance(l, dict) else {}
            on = l.get("on") if isinstance(l.get("on"), dict) else {}
            rows.append({"entity": e, "kind": "lookup", "name": n, "reference": l.get("reference"),
                         "on": ", ".join(f"{a} = {b}" for a, b in on.items()), "columns": C.as_list(l.get("columns")),
                         "missing": l.get("missing"), "why": prop.note(f"{base}.lookups.{n}"),
                         "_required": {"condition": False}})
        fl = w.ans(f"{base}.flags")
        w.dflt(f"{base}.flags")
        for n, cond in (fl.items() if isinstance(fl, dict) else []):
            rows.append({"entity": e, "kind": "flag", "name": n, "condition": cond, "why": prop.note(f"{base}.flags.{n}"),
                         "_required": {"reference": False, "on": False, "columns": False, "missing": False}})
    w.grid(S_LOOK, "Lookups & flags — enrichment and business flags (optional)",
           "Lookup: bring columns in from a reference entity (its key must be unique). Flag: a yes/no business column "
           "defined once here. Leave this tab empty if you need none.",
           lc, rows, example={"entity": "example", "kind": "lookup", "name": "pu", "reference": "zones",
                              "on": "pu_zone = location_id", "columns": "borough", "missing": "keep_null"},
           row_choices={"_spare": 3}, freeze="C")

    # ---------------- KPIs (one column per KPI)
    _kpi_sheet(w, metrics)

    # ---------------- Golden values
    gc = [Col("kpi", "KPI", width=22, hint="PO ★ · a KPI name from the KPIs tab"),
          Col("key", "Key (grain values)", width=46, hint="col=value; col=value — every grain column, dates as YYYY-MM-DD"),
          Col("value", "Value you computed by hand", width=16, hint="number, or null for 'blank'"),
          Col("how", "How you worked it out", width=40, wrap=True, hint="optional — helps the next reviewer")]
    rows = []
    for m, spec in metrics.items():
        gv = C.get_path(A, f"metrics.{m}.golden_values")
        if isinstance(gv, list) and gv:
            for g in gv:
                key = (g or {}).get("key") or {}
                rows.append({"kpi": m, "key": "; ".join(f"{k}={v}" for k, v in key.items()),
                             "value": (g or {}).get("value"), "_required": {"how": False}})
        else:
            for _ in range(3):
                rows.append({"kpi": m, "key": "", "value": "", "_required": {"how": False, "key": False, "value": False}})
    w.grid(S_GOLD, "Golden values ★ — at least 3 numbers per KPI, worked out by hand",
           "No defaults on purpose: these are the only check that catches a formula that runs but means something else. "
           "Work each number out yourself from the raw files and write it here; every build is checked against them. "
           "Leave a KPI's rows blank to decide later (fast profile only): its dashboard works, but nothing can be "
           "released to consumers until they are filled.",
           gc, rows, example={"kpi": "example", "key": "order_date=2024-03-01; region=North", "value": "1250.50",
                              "how": "sum of price in the raw file for that day and region"},
           row_choices={"_spare": 3}, freeze="B")

    # ---------------- Dashboard
    w.qa(S_DASH, "Dashboard", "The dashboard's title, audience and freshness clock (PO).",
         [("serve.title", "", "", None), ("serve.audience", "", "", None), ("serve.runtime", "", "", None),
          ("serve.reference_clock", "", "", None), ("serve.stale_after_hours", "", "", None),
          ("serve.filters", "", "", None), ("serve.port", "", "", None), ("serve.auto_refresh_seconds", "", "", None)])
    vc = [Col("kind", "Kind", width=8, choices=["tile", "chart"], required=False, hint="tile / chart"),
          Col("name", "Name", width=22, required=False, hint="short name (letters, digits, _)"),
          Col("metric", "KPI", "serve.kpis.*.metric", width=22),
          Col("label", "Label / title", "serve.kpis.*.label", width=30, hint="PO · optional · shown on the tile or chart"),
          Col("type", "Chart type", "serve.charts.*.type", width=10),
          Col("x", "X axis", "serve.charts.*.x", width=24),
          Col("color", "Split by", "serve.charts.*.color", width=22)]
    rows = []
    for n, t in (w.ans("serve.kpis") or {}).items() if isinstance(w.ans("serve.kpis"), dict) else []:
        rows.append({"kind": "tile", "name": n, "metric": (t or {}).get("metric"), "label": (t or {}).get("label"),
                     "_required": {"type": False, "x": False, "color": False}})
    for n, ch in (w.ans("serve.charts") or {}).items() if isinstance(w.ans("serve.charts"), dict) else []:
        ch = ch or {}
        rows.append({"kind": "chart", "name": n, "metric": ch.get("metric"), "label": ch.get("title"),
                     "type": ch.get("type"), "x": ch.get("x"), "color": ch.get("color")})
    w.remember("serve.kpis", C.get_path(prop.doc, "serve.kpis"))
    w.remember("serve.charts", C.get_path(prop.doc, "serve.charts"))
    w.grid(S_VIS, "Dashboard visuals — KPI tiles and charts",
           "Tile: one headline number per KPI (ratios are recomputed from numerator and denominator, never averaged). "
           "Chart: a line or bar of a KPI over one of its grain columns, optionally split by another grain column.",
           vc, rows, row_choices={"_spare": 3}, freeze="C")

    # ---------------- After first load (gate B) — only once bronze exists
    _gate_b_sheet(w, ents)

    _fill_start(w)
    _meta_sheet(w)
    out = out or project.loc("workbook_out") / f"{C.get_path(A, 'project.name') or 'warehouse'}-intake.xlsx"
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    w.wb.save(out)
    audit(project, "intake.workbook", file=str(out.name), questions=sum(v[0] for v in w.counts.values()))
    return out, {"counts": w.counts, "analysis": prop.analysis}


def _describe(info: dict, rows) -> str:
    from . import analyze as AN
    return AN.describe(info, rows or 0) if info else "no sample analysed"


def _source_entity_map(A: dict) -> dict:
    out = {}
    for e, spec in (C.get_path(A, "silver.entities") or {}).items():
        spec = spec if isinstance(spec, dict) else {}
        mapping = spec.get("mapping") if isinstance(spec.get("mapping"), dict) else {}
        for s in C.as_list(spec.get("sources")):
            m = mapping.get(s) if isinstance(mapping.get(s), dict) else {}
            rev = {str(v): k for k, v in m.items()}
            for c in (C.get_path(A, f"sources.{s}.schema.columns") or {}):
                out[(s, c)] = (e, rev.get(c, PR.canon(c)))
    return out


def _null_evidence(A: dict, analysis: dict, e: str, cn: str) -> str:
    spec = C.get_path(A, f"silver.entities.{e}") or {}
    mapping = spec.get("mapping") if isinstance(spec.get("mapping"), dict) else {}
    out = []
    for s in C.as_list(spec.get("sources")):
        m = mapping.get(s) if isinstance(mapping.get(s), dict) else {}
        src_col = m.get(cn, cn) if m else cn
        an = analysis.get(s) or {}
        info = (an.get("columns") or {}).get(src_col)
        if info and an.get("rows"):
            pct = 100.0 * info["empty"] / an["rows"]
            out.append(f"{info['empty']:,} of {an['rows']:,} empty ({pct:.1f}%)" + (f" in {s}" if len(C.as_list(spec.get('sources'))) > 1 else ""))
    return "; ".join(out)


KPI_FIELDS = ["entity", "plain_definition", "metric_type", "numerator.measure", "numerator.filter",
              "denominator.measure", "denominator.include.column", "denominator.include.values",
              "denominator.exclude.values", "filters", "zero_denominator", "population", "grain", "time_grain",
              "date_basis", "unit", "precision", "additivity", "valid_range.min", "valid_range.max", "cadence",
              "dispute_owner", "version", "effective_from", "reconciliation.control", "reconciliation.excluded_by",
              "reconciliation.max_excluded_pct", "target_table"]


NEW_KPI_DEFAULTS = {
    "numerator.filter": lambda w: "none", "filters": lambda w: "none", "population": lambda w: "NA",
    "unit": lambda w: "count", "precision": lambda w: "0", "valid_range.min": lambda w: "0",
    "valid_range.max": lambda w: "none", "cadence": lambda w: "daily", "version": lambda w: "1.0.0",
    "effective_from": lambda w: date.today().isoformat(), "reconciliation.excluded_by": lambda w: "none",
    "reconciliation.max_excluded_pct": lambda w: "0", "time_grain": lambda w: "day",
    "dispute_owner": lambda w: (_role_holders(w.answers.get("people") or {}).get("PO") or [None])[0],
}


def _kpi_sheet(w: Writer, metrics: dict) -> None:
    from openpyxl.utils import get_column_letter
    ws = w.sheet(S_KPI, "KPIs ★ — one metric card per column",
                 "Each column is one KPI's card — the definition every number is built and checked from. Fields marked ★ "
                 "are the product owner's decision. To add a KPI, type its name (letters, digits, _) in an empty header "
                 "cell and fill the column; to drop one, clear its name. Golden values go on the next tab.")
    names = list(metrics) + ["", ""]
    first_kpi = 6
    heads = ["Field", "Question", "Owner", "Need", "Allowed answers"] + [n or "(new KPI name)" for n in names]
    widths = [24, 50, 9, 13, 28] + [28] * len(names)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=10)
    for i, (h, wd) in enumerate(zip(heads, widths), start=1):
        c = ws.cell(row=HDR, column=i, value=h)
        c.font = w.font(bold=True, color="FFFFFF")
        c.fill = w.fill(FILL_HDR)
        c.border = w.border
        c.alignment = w.Align(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = wd
        hint = ("do not change", "what is asked", "role", "", "") [i - 1] if i < first_kpi else "KPI name (type over to add)"
        t = ws.cell(row=HINT, column=i, value=hint)
        t.font = w.font(italic=True, size=8, color=GREY_TXT)
        t.fill = w.fill(FILL_HINT)
        t.border = w.border
        if i >= first_kpi:
            c.font = w.font(bold=True, color="FFFFFF") if names[i - first_kpi] else w.font(italic=True, color="FFFFFF")
    dv: dict[tuple, list[str]] = {}
    for r, rel in enumerate(KPI_FIELDS, start=FIRST):
        f = _field(f"metrics.*.{rel}")
        vals = [rel, f.prompt if f else rel, _owner(f), _need(f), _allowed(f)]
        for i, v in enumerate(vals, start=1):
            w.put(ws, r, i, v, kind="info", wrap=i in (2, 5))
        for j, m in enumerate(names):
            col = first_kpi + j
            path = f"metrics.{m}.{rel}" if m else ""
            v = w.ans(path) if m else NEW_KPI_DEFAULTS.get(rel, lambda w: None)(w)
            if m:
                w.dflt(path)
            req = bool(m) and f is not None and (
                f.level in ("M", "M*") or (f.level in ("C", "C*") and K.eval_condition(f.condition, w.answers, [m])))
            w.put(ws, r, col, v, kind="answer", required=req)
            ch = _dv_choices(f)
            if ch:
                dv.setdefault(tuple(ch), []).append(f"{get_column_letter(col)}{r}")
            if m:
                w.counts[S_KPI][0] += 1
                if req and _fmt(v) == "":
                    w.counts[S_KPI][1] += 1
        ws.row_dimensions[r].height = 30
    # KPI-name header cells are inputs too (type a name to add a KPI)
    for j, m in enumerate(names):
        c = ws.cell(row=HDR, column=first_kpi + j)
        c.value = m or None
        c.fill = w.fill("2F5597")
        c.number_format = "@"
    w.validate(ws, dv)
    ws.freeze_panes = ws.cell(row=FIRST, column=3)


def _gate_b_sheet(w: Writer, ents: dict) -> None:
    from . import silver
    rows = []
    project = w.project
    for e in ents:
        try:
            amb = silver.ambiguity_counts(project, e)
            # re-render the read-back only when the owner has not confirmed THESE rules yet:
            # writing it again re-stamps it and would cancel a valid confirmation
            rb = silver.rule_readback(project, e, write=False)
            if rb.get("rules") and not rb.get("readback_current"):
                rb = silver.rule_readback(project, e, write=True)
        except Exception:
            continue
        for col, n in amb.items():
            rows.append({"entity": e, "question": "Ambiguous dates", "item": col,
                         "evidence": f"{n:,} values read as different dates under two declared formats",
                         "decision": w.ans(f"silver.entities.{e}.date_ambiguity.{col}") or "reject",
                         "_choices": {"decision": ["prefer_first", "reject"]},
                         "_pattern": {"decision": "silver.entities.*.date_ambiguity.*"}})
            w.remember(f"silver.entities.{e}.date_ambiguity.{col}", "reject")
        if rb.get("rules"):
            ev = "\n".join(f"{r['kind']} '{r['name']}': {r['matches']:,} rows → {r['action']}" for r in rb["rows"])
            if rb.get("conflicts"):
                ev += "\nCONFLICTS: " + "; ".join(rb["conflicts"])
            rows.append({"entity": e, "question": "Rule read-back", "item": "readback_confirmed",
                         "evidence": ev,
                         "decision": w.ans(f"silver.entities.{e}.readback_confirmed") if rb.get("readback_current") else None,
                         "_choices": {"decision": ["yes", "no"]},
                         "_pattern": {"decision": "silver.entities.*.readback_confirmed"},
                         "_required": {"decision": True}})
        doc = project.document()
        from . import validators as V
        for cand in V.profile_null_candidates(doc, e):
            if cand in (C.get_path(w.answers, f"silver.entities.{e}.structural_nulls") or {}):
                continue
            rows.append({"entity": e, "question": "Never filled", "item": cand,
                         "evidence": "the bronze profile shows this ~100% empty", "decision": "defect",
                         "_choices": {"decision": ["structural", "defect"]},
                         "_pattern": {"decision": "silver.entities.*.structural_nulls.*"}})
            w.remember(f"silver.entities.{e}.structural_nulls.{cand}", "defect")
    if not rows:
        return
    w.grid(S_GATEB, "After first load ★ — questions only the loaded data can raise",
           "Ambiguous dates: values like 03/04/2024 that two declared formats read differently. Rule read-back: each rule as "
           "it will run, with how many loaded rows it matches — confirm (yes) that they mean what you intended; no default "
           "on purpose. Never filled: columns empty in the loaded data.",
           [Col("entity", "Entity", kind="key", width=15), Col("question", "Question", kind="key", width=16),
            Col("item", "Column / item", kind="key", width=28), Col("evidence", "What the data shows (counts only)",
                                                                   kind="info", width=70, wrap=True),
            Col("decision", "Decision", width=14, hint="SME ★ · see the list")], rows, freeze="D")


def _start_sheet(w: Writer) -> None:
    w.wb.create_sheet(S_START, 0)


def _fill_start(w: Writer) -> None:
    ws = w.wb[S_START]
    ws.sheet_view.showGridLines = False
    A = w.answers
    name = C.get_path(A, "project.name") or "warehouse"
    lines = [
        (f"Intake workbook — {name}", "title"),
        (f"Generated {now_iso()} by dwh_core {VERSION}. Every question every layer needs, before anything is built.", "sub"),
        ("", None),
        ("How to fill it", "h"),
        ("1. Go through the tabs left to right. Each answer cell is pre-filled with a default (blue text on yellow).", None),
        ("2. Keep a default as it is, or type over it. Use the drop-down where a cell has one.", None),
        ("3. Orange cells are required and have no default — they must be filled (golden values may wait, see that tab).", None),
        ("4. NA = not applicable (only where the hint says optional). none = there are none. pending = an owner will "
         "decide later (only where offered; blocks release to consumers).", None),
        ("5. ★ marks an owner decision. Keeping its default records it as that owner's own answer, so read it first.", None),
        ("6. Grey cells are information (counts from the sample files — never data values). Do not change the "
         "'Question ID' and 'Field' columns or the header rows.", None),
        ("7. Save and send the file back. It is checked as a whole: if anything is missing or inconsistent, nothing is "
         "recorded and you get it back with every problem marked on its cell and listed on an 'Issues' tab.", None),
        ("", None),
        ("Colour key", "h"),
        ("Default you may keep or change", "in"), ("Required — no default, must be filled", "req"),
        ("Information only", "info"), ("Example row — ignored", "ex"),
        ("", None),
        ("Tabs", "h"),
    ]
    r = 1
    for text, style in lines:
        c = ws.cell(row=r, column=1, value=text)
        if style == "title":
            c.font = w.font(size=16, bold=True, color=NAVY)
        elif style == "sub":
            c.font = w.font(color=GREY_TXT)
        elif style == "h":
            c.font = w.font(size=12, bold=True, color=NAVY)
        elif style in ("in", "req", "info", "ex"):
            c.font = w.font(color=BLUE if style == "in" else GREY_TXT if style == "info" else "000000",
                            italic=style == "ex")
            c.fill = w.fill({"in": FILL_IN, "req": FILL_REQ, "info": FILL_INFO, "ex": "FFFFFF"}[style])
            c.border = w.border
        else:
            c.font = w.font()
        r += 1
    owners = {S_PEOPLE: "DE", S_PROJECT: "DE", S_GOV: "GOV ★", S_POLICY: "PO ★", S_SOURCES: "DE", S_SCHEMA: "DE",
              S_CLASS: "GOV ★", S_SILVER: "DE", S_ENT: "DE · SME · PO", S_DQ: "SME (★ where marked)",
              S_LOOK: "DE · SME", S_KPI: "PO ★", S_GOLD: "PO ★", S_DASH: "PO", S_VIS: "PO", S_GATEB: "SME ★"}
    about = {S_PEOPLE: "who answers and approves", S_PROJECT: "name, profile, who answered each role",
             S_GOV: "compliance, what the assistant may see, row security", S_POLICY: "time zone, calendar, money",
             S_SOURCES: "how each feed lands (bronze)", S_SCHEMA: "column types and keys — mandatory",
             S_CLASS: "how each column is protected", S_SILVER: "where each column goes in silver",
             S_ENT: "keys, duplicates, merge per batch", S_DQ: "rules, empty values, never-filled columns",
             S_LOOK: "enrichment and business flags", S_KPI: "metric cards", S_GOLD: "hand-computed checks",
             S_DASH: "title, audience, freshness", S_VIS: "tiles and charts",
             S_GATEB: "ambiguous dates and rule read-back (after the first load)"}
    hdr = ["Tab", "What it asks", "Who answers", "Answers", "Required, no default"]
    for i, h in enumerate(hdr, start=1):
        c = ws.cell(row=r, column=i, value=h)
        c.font = w.font(bold=True, color="FFFFFF")
        c.fill = w.fill(FILL_HDR)
        c.border = w.border
    r += 1
    for tab in [s for s in w.wb.sheetnames if s not in (S_START, S_META)]:
        n, req = w.counts.get(tab, [0, 0])
        for i, v in enumerate([tab, about.get(tab, ""), owners.get(tab, ""), n, req], start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = w.font(bold=(i == 1))
            c.border = w.border
            if i == 1:
                c.hyperlink = f"#'{tab}'!A1"
                c.font = w.font(bold=True, color=BLUE, underline="single")
        r += 1
    for col, wd in zip("ABCDE", (34, 52, 20, 10, 18)):
        ws.column_dimensions[col].width = wd


def _meta_sheet(w: Writer) -> None:
    ws = w.wb.create_sheet(S_META)
    wid = uuid.uuid4().hex
    rows = [("workbook", "dwh intake workbook"), ("version", WB_VERSION), ("kernel", VERSION),
            ("project", C.get_path(w.answers, "project.name") or ""), ("generated_at", now_iso()), ("id", wid)]
    for r, (k, v) in enumerate(rows, start=1):
        ws.cell(row=r, column=1, value=k)
        ws.cell(row=r, column=2, value=v)
    ws.cell(row=8, column=1, value="default_path")
    ws.cell(row=8, column=2, value="default_value_json")
    r = 9
    for p, v in sorted(w.defaults.items()):
        ws.cell(row=r, column=1, value=p)
        ws.cell(row=r, column=2, value=json.dumps(C.to_strings(v), ensure_ascii=False))
        r += 1
    ws.sheet_state = "hidden"


# ================================================================ reading
@dataclass
class Issue:
    sheet: str
    cell: str
    path: str
    problem: str
    fix: str = ""


@dataclass
class Parsed:
    doc: dict = field(default_factory=dict)
    cells: dict = field(default_factory=dict)       # path -> (sheet, coord)
    anchors: dict = field(default_factory=dict)     # path prefix -> (sheet, coord)
    issues: list = field(default_factory=list)
    answered_by: dict = field(default_factory=dict)  # role -> person
    defaults: dict = field(default_factory=dict)
    pending: list = field(default_factory=list)     # (path, role)
    gate_b: bool = False
    meta: dict = field(default_factory=dict)


class Reader:
    def __init__(self, path: Path):
        import openpyxl
        self.path = Path(path)
        try:
            self.wb = openpyxl.load_workbook(self.path, data_only=True)
        except Exception as e:
            raise WorkbookError(f"cannot open {self.path.name} as an .xlsx workbook: {e}") from None
        self.P = Parsed()

    def issue(self, sheet, coord, path, problem, fix=""):
        self.P.issues.append(Issue(sheet, coord, path, problem, fix))

    def ws(self, name: str, required: bool = True):
        if name in self.wb.sheetnames:
            return self.wb[name]
        if required:
            self.issue(name, "", "", f"tab '{name}' is missing", "send back the whole workbook, with every tab")
        return None

    def table(self, name: str, labels: dict[str, str], required: bool = True) -> list[tuple[int, dict, dict]]:
        """Rows of a grid tab: (row number, {key: text}, {key: coordinate})."""
        from openpyxl.utils import get_column_letter
        ws = self.ws(name, required)
        if ws is None:
            return []
        hdr_row, pos = None, {}
        for r in range(1, 12):
            vals = {_norm(ws.cell(row=r, column=c).value): c for c in range(1, ws.max_column + 1)}
            if _norm(next(iter(labels.values()))) in vals:
                hdr_row = r
                for k, lab in labels.items():
                    if _norm(lab) in vals:
                        pos[k] = vals[_norm(lab)]
                break
        if hdr_row is None:
            self.issue(name, "A4", "", "header row not found", "do not change the header rows")
            return []
        for k, lab in labels.items():
            if k not in pos:
                self.issue(name, f"A{hdr_row}", "", f"column '{lab}' is missing", "do not delete or rename columns")
        out = []
        for r in range(hdr_row + 2, ws.max_row + 1):
            row, coords = {k: "" for k in labels}, {k: f"A{r}" for k in labels}
            for k, c in pos.items():
                row[k] = _cell(ws.cell(row=r, column=c).value)
                coords[k] = f"{get_column_letter(c)}{r}"
            if not any(row.values()):
                continue
            first = row.get(next(iter(labels)), "")
            if str(first).lower().startswith("example"):
                continue
            out.append((r, row, coords))
        return out

    def qa(self, name: str) -> dict[str, tuple[str, str]]:
        rows = self.table(name, {"id": "Question ID", "answer": "Your answer"})
        return {row["id"]: (row["answer"], coords["answer"]) for _, row, coords in rows if row.get("id")}

    # ---- value placement
    def put(self, path: str, text: str, sheet: str, coord: str, f: K.Field | None = None, as_list: bool | None = None):
        self.P.cells[path] = (sheet, coord)
        if text == "":
            return
        if f is None:
            f = _field(path)
        low = text.strip().lower()
        if f is not None and low == "pending":
            if f.human_owned and f.pending_ok:
                self.P.pending.append((path, f.owner))
            else:
                self.issue(sheet, coord, path, "'pending' is not possible here — this answer is needed to build",
                           "give the answer now")
            return
        lst = as_list if as_list is not None else (f is not None and f.type in ("list", "columns", "strftime_list"))
        if lst and low not in ("na", "none"):
            value: Any = _split(text)
        elif f is not None and f.type == "enum" and low in {c.lower() for c in f.choices}:
            value = low if text not in f.choices else text
        elif low in ("na",):
            value = "NA"
        elif low == "none":
            value = "none"
        else:
            value = text
        C.set_path(self.P.doc, path, value)

    def anchor(self, prefix: str, sheet: str, coord: str):
        self.P.anchors.setdefault(prefix, (sheet, coord))


def parse(path: Path) -> Parsed:
    R = Reader(path)
    P = R.P
    meta = R.ws(S_META)
    if meta is not None:
        for r in range(1, 8):
            k, v = meta.cell(row=r, column=1).value, meta.cell(row=r, column=2).value
            if k:
                P.meta[str(k)] = _cell(v)
        for r in range(9, meta.max_row + 1):
            k, v = meta.cell(row=r, column=1).value, meta.cell(row=r, column=2).value
            if k:
                try:
                    P.defaults[str(k)] = json.loads(v) if v else None
                except (TypeError, ValueError):
                    pass
    if P.meta.get("workbook") != "dwh intake workbook":
        raise WorkbookError("this is not a dwh intake workbook (its hidden _meta tab is missing or changed)")

    # ---- People
    for r, row, co in R.table(S_PEOPLE, {"id": "Person id", "name": "Full name", "email": "Email", "roles": "Roles"}):
        pid = row["id"].strip()
        if not pid:
            R.issue(S_PEOPLE, co["id"], "people", "a person needs an id", "short id, e.g. sandeep")
            continue
        if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_.-]*", pid):
            R.issue(S_PEOPLE, co["id"], f"people.{pid}", "person ids use letters, digits, _ . -", "")
            continue
        R.anchor(f"people.{pid}", S_PEOPLE, co["id"])
        R.put(f"people.{pid}.name", row["name"], S_PEOPLE, co["name"])
        R.put(f"people.{pid}.email", row["email"], S_PEOPLE, co["email"])
        roles = [x.strip().upper() for x in _split(row["roles"])]
        P.cells[f"people.{pid}.roles"] = (S_PEOPLE, co["roles"])
        if roles:
            C.set_path(P.doc, f"people.{pid}.roles", roles)
    R.anchor("people", S_PEOPLE, f"A{FIRST}")

    # ---- Q&A tabs
    for sheet in (S_PROJECT, S_GOV, S_POLICY, S_DASH):
        for qid, (ans, coord) in R.qa(sheet).items():
            if qid.startswith("answered_by."):
                role = qid.split(".", 1)[1]
                P.cells[qid] = (sheet, coord)
                if ans:
                    P.answered_by[role] = ans.strip()
                continue
            f = _field(qid)
            if f is None:
                R.issue(sheet, coord, qid, "unknown question id", "do not change the Question ID column")
                continue
            R.put(qid, ans, sheet, coord, f)

    # ---- Sources
    src_labels = {"source": "Source", "role": "Role", "connector": "Connector", "location": "Location",
                  "format": "Format", "load_type": "Load type", "redelivery_policy": "Re-delivered batch",
                  "drift_policy": "Unexpected new columns", "freshness.basis": "Freshness basis",
                  "freshness.token_format": "Batch id format", "freshness.column": "Freshness column",
                  "freshness.formats": "Freshness column formats", "freshness.window_days": "Max age (days)",
                  "volume.min_rows": "Min rows", "volume.max_rows": "Max rows",
                  "volume.threshold_source": "Row bounds from", "retry_attempts": "Retries",
                  "credentials_env": "Credentials env var", "profile_split_by": "Profile split by",
                  "csv.delimiter": "CSV delimiter", "csv.header": "CSV header row", "csv.quote": "CSV quote",
                  "csv.escape": "CSV escape", "csv.encoding": "CSV encoding", "csv.null_tokens": "CSV null tokens",
                  "xlsx.sheet": "Excel sheet", "xlsx.header_row": "Excel header row"}
    sources: list[str] = []
    for r, row, co in R.table(S_SOURCES, src_labels):
        s = row["source"].strip()
        if not s:
            R.issue(S_SOURCES, co["source"], "sources", "a source needs a name", "")
            continue
        if s in sources:
            R.issue(S_SOURCES, co["source"], f"sources.{s}", f"source '{s}' is listed twice", "")
            continue
        sources.append(s)
        R.anchor(f"sources.{s}", S_SOURCES, co["source"])
        fmt = row.get("format", "").lower()
        for k in src_labels:
            if k == "source":
                continue
            if k.startswith("csv.") and fmt != "csv" and not row.get(k):
                continue
            if k.startswith("xlsx.") and fmt != "xlsx":
                continue
            R.put(f"sources.{s}.{k}", row.get(k, ""), S_SOURCES, co[k], _field(f"sources.*.{k}"))
    R.anchor("sources", S_SOURCES, f"A{FIRST}")

    # ---- Schema
    for r, row, co in R.table(S_SCHEMA, {"source": "Source", "column": "Column", "type": "Type",
                                          "required": "Required in every file", "key": "Part of the record key",
                                          "allowed_values": "Allowed values"}):
        s, c = row["source"].strip(), row["column"].strip()
        if s not in sources:
            R.issue(S_SCHEMA, co["source"], "sources", f"unknown source '{s}'", "use a name from the Sources tab")
            continue
        if not c:
            R.issue(S_SCHEMA, co["column"], f"sources.{s}.schema.columns", "a column needs a name", "")
            continue
        base = f"sources.{s}.schema.columns.{c}"
        R.anchor(base, S_SCHEMA, co["column"])
        R.anchor(f"sources.{s}.schema", S_SCHEMA, co["source"])
        for k in ("type", "required", "key", "allowed_values"):
            R.put(f"{base}.{k}", row[k], S_SCHEMA, co[k], _field(f"sources.*.schema.columns.*.{k}"))
            if k == "type" and row[k]:
                C.set_path(P.doc, f"{base}.type", row[k].upper().replace(" ", "") if "(" in row[k] else row[k].upper())
    for s in sources:
        P.anchors.setdefault(f"sources.{s}.schema", (S_SCHEMA, f"A{FIRST}"))

    # ---- Silver model
    ent_sources: dict[str, list[str]] = {}
    ent_cols: dict[str, dict[str, dict]] = {}
    ent_map: dict[str, dict[str, dict[str, str]]] = {}
    col_origin: dict[tuple, tuple[str, str]] = {}
    sm_labels = {"source": "Source", "column": "Source column", "entity": "Entity", "silver_column": "Silver column",
                 "expression": "Expression (optional)", "type": "Silver type", "formats": "Date/time formats",
                 "source_timezone": "Recorded in time zone", "drop": "Not carried — reason"}
    for r, row, co in R.table(S_SILVER, sm_labels):
        s, c, e = row["source"].strip(), row["column"].strip(), row["entity"].strip()
        if s not in sources:
            R.issue(S_SILVER, co["source"], "silver.entities", f"unknown source '{s}'", "use a name from the Sources tab")
            continue
        if not e:
            R.issue(S_SILVER, co["entity"], "silver.entities", "which entity does this column go to?",
                    "fill Entity (or write a reason under 'Not carried')")
            continue
        if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", e):
            R.issue(S_SILVER, co["entity"], f"silver.entities.{e}", "entity names use letters, digits and _", "")
            continue
        ent_sources.setdefault(e, [])
        if s not in ent_sources[e]:
            ent_sources[e].append(s)
        R.anchor(f"silver.entities.{e}.columns", S_SILVER, co["entity"])
        R.anchor(f"silver.entities.{e}.mapping", S_SILVER, co["entity"])
        if row["drop"]:
            if not c:
                R.issue(S_SILVER, co["drop"], f"silver.entities.{e}.drop_columns", "a dropped row needs its source column", "")
                continue
            C.set_path(P.doc, f"silver.entities.{e}.drop_columns.{c}", row["drop"])
            P.cells[f"silver.entities.{e}.drop_columns.{c}"] = (S_SILVER, co["drop"])
            R.anchor(f"silver.entities.{e}.drop_columns", S_SILVER, co["drop"])
            continue
        cn = row["silver_column"].strip() or (PR.canon(c) if c else "")
        if not cn:
            R.issue(S_SILVER, co["silver_column"], f"silver.entities.{e}.columns", "a silver column name is needed", "")
            continue
        expr = row["expression"].strip()
        if not c and not expr:
            R.issue(S_SILVER, co["expression"], f"silver.entities.{e}.columns.{cn}",
                    "a row without a source column needs an expression", "")
            continue
        if c:
            col_origin[(s, c)] = (e, cn)
        cols = ent_cols.setdefault(e, {})
        spec = cols.setdefault(cn, {})
        base = f"silver.entities.{e}.columns.{cn}"
        for k in ("type", "formats", "source_timezone"):
            txt = row[k]
            if not txt:
                P.cells.setdefault(f"{base}.{k}", (S_SILVER, co[k]))
                continue
            val = _split(txt) if k == "formats" else (txt.upper() if k == "type" else txt)
            if k in spec and spec[k] != val:
                R.issue(S_SILVER, co[k], f"{base}.{k}",
                        f"'{cn}' is given different {k} values on different rows", "make them the same")
                continue
            spec[k] = val
            P.cells[f"{base}.{k}"] = (S_SILVER, co[k])
        ent_map.setdefault(e, {}).setdefault(s, {})[cn] = expr or c
        P.cells[f"silver.entities.{e}.mapping.{s}.{cn}"] = (S_SILVER, co["expression"] if expr else co["silver_column"])
    for e, cols in ent_cols.items():
        C.set_path(P.doc, f"silver.entities.{e}.sources", ent_sources[e])
        C.set_path(P.doc, f"silver.entities.{e}.columns", cols)
        m = ent_map.get(e, {})
        needs = len(ent_sources[e]) > 1 or any(src != cn for sm in m.values() for cn, src in sm.items())
        if needs:
            full = {}
            for s in ent_sources[e]:
                full[s] = {cn: m.get(s, {}).get(cn, "NULL") for cn in cols}
            C.set_path(P.doc, f"silver.entities.{e}.mapping", full)
    for e in ent_sources:
        if e not in ent_cols:
            C.set_path(P.doc, f"silver.entities.{e}.sources", ent_sources[e])

    # ---- Entities
    ent_labels = {"entity": "Entity", "natural_key": "Record key", "tiebreaker": "Tiebreaker",
                  "merge.strategy": "Merge strategy", "merge.version_column": "Version column",
                  "merge.partition_column": "Partition column", "merge.op_column": "CDC op column",
                  "merge.op_codes.insert": "CDC insert code", "merge.op_codes.update": "CDC update code",
                  "merge.op_codes.delete": "CDC delete code", "merge.sequence_column": "CDC sequence column",
                  "rejected_survivor_policy": "If the newest version breaks a rule",
                  "dead_letter_tolerance_pct": "Max % of a batch rejected", "tolerance_source": "Tolerance from",
                  "target_table": "Silver table name"}
    seen_ents = set()
    for r, row, co in R.table(S_ENT, ent_labels):
        e = row["entity"].strip()
        if e not in ent_sources:
            R.issue(S_ENT, co["entity"], f"silver.entities.{e}", f"entity '{e}' has no columns on the Silver model tab",
                    "map columns to it there, or delete this row")
            continue
        seen_ents.add(e)
        R.anchor(f"silver.entities.{e}", S_ENT, co["entity"])
        for k in ent_labels:
            if k == "entity":
                continue
            R.put(f"silver.entities.{e}.{k}", row.get(k, ""), S_ENT, co[k], _field(f"silver.entities.*.{k}"))
    for e in ent_sources:
        if e not in seen_ents:
            R.issue(S_ENT, f"A{FIRST}", f"silver.entities.{e}", f"entity '{e}' needs a row on the Entities tab", "")

    # ---- Classification (+ masking)
    for r, row, co in R.table(S_CLASS, {"source": "Source", "column": "Column", "classification": "Classification",
                                         "masking": "Handling in silver"}):
        s, c = row["source"].strip(), row["column"].strip()
        if s not in sources:
            R.issue(S_CLASS, co["source"], "sources", f"unknown source '{s}'", "")
            continue
        base = f"sources.{s}.schema.columns.{c}"
        R.put(f"{base}.classification", row["classification"], S_CLASS, co["classification"],
              _field("sources.*.schema.columns.*.classification"))
        e_cn = col_origin.get((s, c))
        if e_cn:
            mp = f"silver.entities.{e_cn[0]}.masking.{e_cn[1]}"
            P.cells[mp] = (S_CLASS, co["masking"])
            if row["masking"]:
                prev = C.get_path(P.doc, mp)
                if prev is not None and str(prev).lower() != row["masking"].lower():
                    R.issue(S_CLASS, co["masking"], mp, f"silver column {e_cn[1]} gets different handling from two source columns", "")
                R.put(mp, row["masking"], S_CLASS, co["masking"], _field("silver.entities.*.masking.*"))
        R.anchor(f"sources.{s}.schema.columns.{c}.classification", S_CLASS, co["classification"])

    # ---- Data quality
    dq_labels = {"entity": "Entity", "check": "Check", "name": "Column / rule name", "condition": "Condition (SQL)",
                 "reason": "Reason", "decision": "Decision"}
    said_none: dict[tuple, tuple] = {}
    applied: dict[tuple, int] = {}
    for r, row, co in R.table(S_DQ, dq_labels):
        e, chk, dec = row["entity"].strip(), row["check"].strip(), row["decision"].strip()
        if e not in ent_sources:
            R.issue(S_DQ, co["entity"], "silver.entities", f"unknown entity '{e}'", "use a name from the Entities tab")
            continue
        base = f"silver.entities.{e}"
        kind = {CHECK_NULL.lower(): "null", CHECK_HARD.lower(): "hard_rejects", CHECK_ANOM.lower(): "valid_anomalies",
                CHECK_STRUCT.lower(): "struct"}.get(chk.lower())
        if kind is None:
            R.issue(S_DQ, co["check"], base, f"unknown check '{chk}'",
                    f"use {CHECK_HARD}, {CHECK_ANOM}, {CHECK_NULL} or {CHECK_STRUCT}")
            continue
        if kind == "null":
            c = row["name"].strip()
            R.anchor(f"{base}.null_policy", S_DQ, co["decision"])
            R.put(f"{base}.null_policy.{c}", dec, S_DQ, co["decision"], _field("silver.entities.*.null_policy.*"))
            continue
        if kind == "struct":
            cand = row["name"].strip()
            R.anchor(f"{base}.structural_nulls", S_DQ, co["decision"])
            _put_struct(R, base, cand, dec, S_DQ, co["decision"])
            continue
        R.anchor(f"{base}.{kind}", S_DQ, co["decision"])
        low = dec.lower()
        if low == "none":
            said_none[(e, kind)] = (S_DQ, co["decision"])
            continue
        if low == "skip":  # considered and not wanted: with no applied rule the answer is 'none'
            said_none.setdefault((e, kind), (S_DQ, co["decision"]))
            continue
        if low != "apply":
            R.issue(S_DQ, co["decision"], f"{base}.{kind}", f"decision must be apply, skip or none (got '{dec}')", "")
            continue
        name = row["name"].strip()
        if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", name or ""):
            R.issue(S_DQ, co["name"], f"{base}.{kind}", "a rule needs a short name (letters, digits, _)", "")
            continue
        applied[(e, kind)] = applied.get((e, kind), 0) + 1
        R.put(f"{base}.{kind}.{name}.predicate", row["condition"], S_DQ, co["condition"], _field(f"silver.entities.*.{kind}.*.predicate"))
        R.put(f"{base}.{kind}.{name}.reason", row["reason"], S_DQ, co["reason"], _field(f"silver.entities.*.{kind}.*.reason"))
        P.cells[f"{base}.{kind}.{name}"] = (S_DQ, co["name"])
    for (e, kind), cell in said_none.items():
        if not applied.get((e, kind)):
            C.set_path(P.doc, f"silver.entities.{e}.{kind}", "none")
            P.cells[f"silver.entities.{e}.{kind}"] = cell

    # ---- Lookups & flags
    lk_labels = {"entity": "Entity", "kind": "Kind", "name": "Name", "reference": "Reference entity", "on": "Join on",
                 "columns": "Columns to bring", "missing": "If the code is missing", "condition": "Flag condition (SQL)"}
    has_lk, has_fl = set(), set()
    for r, row, co in R.table(S_LOOK, lk_labels, required=False):
        e, kind, name = row["entity"].strip(), row["kind"].strip().lower(), row["name"].strip()
        if e not in ent_sources:
            R.issue(S_LOOK, co["entity"], "silver.entities", f"unknown entity '{e}'", "")
            continue
        if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", name or ""):
            R.issue(S_LOOK, co["name"], f"silver.entities.{e}", "a short name is needed (letters, digits, _)", "")
            continue
        base = f"silver.entities.{e}"
        if kind == "lookup":
            has_lk.add(e)
            R.anchor(f"{base}.lookups.{name}", S_LOOK, co["name"])
            R.put(f"{base}.lookups.{name}.reference", row["reference"], S_LOOK, co["reference"])
            on = {}
            for pair in _split(row["on"]):
                a, sep, b = pair.partition("=")
                if not sep or not a.strip() or not b.strip():
                    R.issue(S_LOOK, co["on"], f"{base}.lookups.{name}.on", f"'{pair}' is not local = reference", "")
                    continue
                on[a.strip()] = b.strip()
            P.cells[f"{base}.lookups.{name}.on"] = (S_LOOK, co["on"])
            if on:
                C.set_path(P.doc, f"{base}.lookups.{name}.on", on)
            R.put(f"{base}.lookups.{name}.columns", row["columns"], S_LOOK, co["columns"], as_list=True)
            R.put(f"{base}.lookups.{name}.missing", row["missing"], S_LOOK, co["missing"])
        elif kind == "flag":
            has_fl.add(e)
            R.put(f"{base}.flags.{name}", row["condition"], S_LOOK, co["condition"], _field("silver.entities.*.flags.*"))
        else:
            R.issue(S_LOOK, co["kind"], base, "kind must be lookup or flag", "")
    for e in ent_sources:
        if e not in has_lk:
            C.set_path(P.doc, f"silver.entities.{e}.lookups", "NA")
        if e not in has_fl:
            C.set_path(P.doc, f"silver.entities.{e}.flags", "NA")
        if C.get_path(P.doc, f"silver.entities.{e}.drop_columns") is None:
            C.set_path(P.doc, f"silver.entities.{e}.drop_columns", "NA")

    # ---- KPIs
    metrics = _parse_kpis(R)

    # ---- Golden values
    golden: dict[str, list] = {}
    gold_cells: dict[str, tuple] = {}
    for r, row, co in R.table(S_GOLD, {"kpi": "KPI", "key": "Key (grain values)", "value": "Value you computed by hand"}):
        m = row["kpi"].strip()
        if not row["key"] and not row["value"]:
            if m in metrics:
                gold_cells.setdefault(m, (S_GOLD, co["key"]))
            continue  # an empty row (also for a KPI that was removed)
        if m not in metrics:
            R.issue(S_GOLD, co["kpi"], "metrics", f"unknown KPI '{m}'", "use a name from the KPIs tab")
            continue
        gold_cells.setdefault(m, (S_GOLD, co["key"]))
        key = {}
        for part in [p for p in row["key"].split(";") if p.strip()]:
            a, sep, b = part.partition("=")
            if not sep:
                R.issue(S_GOLD, co["key"], f"metrics.{m}.golden_values", f"'{part.strip()}' is not col=value", "")
                continue
            key[a.strip()] = b.strip()
        val = row["value"].strip()
        if not key or val == "":
            R.issue(S_GOLD, co["key"] if not key else co["value"], f"metrics.{m}.golden_values",
                    "a golden value needs both its key and its value", "")
            continue
        if val.lower() == "null":
            val = "null"
        else:
            try:
                C.as_number(val)
            except ValueError:
                R.issue(S_GOLD, co["value"], f"metrics.{m}.golden_values", f"'{val}' is not a number (or null)", "")
                continue
        golden.setdefault(m, []).append({"key": key, "value": val})
    for m in metrics:
        P.anchors.setdefault(f"metrics.{m}.golden_values", gold_cells.get(m, (S_GOLD, f"A{FIRST}")))
        P.cells[f"metrics.{m}.golden_values"] = gold_cells.get(m, (S_GOLD, f"A{FIRST}"))
        if golden.get(m):
            C.set_path(P.doc, f"metrics.{m}.golden_values", golden[m])
        else:
            P.pending.append((f"metrics.{m}.golden_values", "PO"))

    # ---- Dashboard visuals
    for r, row, co in R.table(S_VIS, {"kind": "Kind", "name": "Name", "metric": "KPI", "label": "Label / title",
                                       "type": "Chart type", "x": "X axis", "color": "Split by"}):
        kind, n = row["kind"].strip().lower(), row["name"].strip()
        if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", n or ""):
            R.issue(S_VIS, co["name"], "serve", "a short name is needed (letters, digits, _)", "")
            continue
        if kind == "tile":
            R.anchor(f"serve.kpis.{n}", S_VIS, co["name"])
            R.put(f"serve.kpis.{n}.metric", row["metric"], S_VIS, co["metric"])
            R.put(f"serve.kpis.{n}.label", row["label"] or "NA", S_VIS, co["label"])
        elif kind == "chart":
            R.anchor(f"serve.charts.{n}", S_VIS, co["name"])
            for k, ck in (("metric", "metric"), ("title", "label"), ("type", "type"), ("x", "x"), ("color", "color")):
                txt = row[ck] or ("NA" if k in ("title", "color") else "")
                R.put(f"serve.charts.{n}.{k}", txt, S_VIS, co[ck])
        else:
            R.issue(S_VIS, co["kind"], "serve", "kind must be tile or chart", "")
    P.anchors.setdefault("serve.kpis", (S_VIS, f"A{FIRST}"))
    P.anchors.setdefault("serve.charts", (S_VIS, f"A{FIRST}"))
    P.anchors.setdefault("serve", (S_DASH, f"A{FIRST}"))

    # ---- After first load
    if S_GATEB in R.wb.sheetnames:
        P.gate_b = True
        for r, row, co in R.table(S_GATEB, {"entity": "Entity", "question": "Question", "item": "Column / item",
                                             "decision": "Decision"}):
            e, q, item = row["entity"].strip(), row["question"].strip().lower(), row["item"].strip()
            base = f"silver.entities.{e}"
            if q == "ambiguous dates":
                R.put(f"{base}.date_ambiguity.{item}", row["decision"], S_GATEB, co["decision"])
            elif q == "rule read-back":
                R.put(f"{base}.readback_confirmed", row["decision"], S_GATEB, co["decision"])
            elif q == "never filled":
                _put_struct(R, base, item, row["decision"], S_GATEB, co["decision"])
    _check_roles(R)
    return P


def _put_struct(R: Reader, base: str, cand: str, dec: str, sheet: str, coord: str) -> None:
    path = f"{base}.structural_nulls.{cand}"
    R.P.cells[path] = (sheet, coord)
    if not dec:
        return
    if dec.strip().lower() not in ("structural", "defect"):
        R.issue(sheet, coord, path, f"decision must be structural or defect (got '{dec}')", "")
        return
    sn = C.get_path(R.P.doc, f"{base}.structural_nulls")
    if not isinstance(sn, dict):
        sn = {}
        C.set_path(R.P.doc, f"{base}.structural_nulls", sn)
    sn[cand] = dec.strip().lower()


def _parse_kpis(R: Reader) -> list[str]:
    ws = R.ws(S_KPI)
    if ws is None:
        return []
    rows = {}
    for r in range(FIRST, ws.max_row + 1):
        k = _cell(ws.cell(row=r, column=1).value)
        if k:
            rows[k] = r
    missing = [k for k in KPI_FIELDS if k not in rows]
    if missing:
        R.issue(S_KPI, "A6", "metrics", f"rows missing: {', '.join(missing)}", "do not delete rows of the KPIs tab")
    from openpyxl.utils import get_column_letter
    names = []
    for col in range(6, ws.max_column + 1):
        name = _cell(ws.cell(row=HDR, column=col).value)
        if not name or name.startswith("(new KPI"):
            continue
        letter = get_column_letter(col)
        if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", name):
            R.issue(S_KPI, f"{letter}{HDR}", "metrics", f"KPI name '{name}' must use letters, digits and _", "")
            continue
        if name in names:
            R.issue(S_KPI, f"{letter}{HDR}", "metrics", f"KPI '{name}' appears twice", "")
            continue
        names.append(name)
        R.anchor(f"metrics.{name}", S_KPI, f"{letter}{HDR}")
        for rel, r in rows.items():
            if rel not in KPI_FIELDS:
                continue
            R.put(f"metrics.{name}.{rel}", _cell(ws.cell(row=r, column=col).value), S_KPI, f"{letter}{r}",
                  _field(f"metrics.*.{rel}"))
    R.P.anchors.setdefault("metrics", (S_KPI, "F4"))
    return names


def _check_roles(R: Reader) -> None:
    P = R.P
    ppl = P.doc.get("people") or {}
    for role, who in list(P.answered_by.items()):
        cell = P.cells.get(f"answered_by.{role}", (S_PROJECT, ""))
        if who.lower() in AGENT_IDS:
            R.issue(cell[0], cell[1], f"answered_by.{role}", "the assistant cannot answer for a role",
                    "name the person who did")
        elif who not in ppl:
            R.issue(cell[0], cell[1], f"answered_by.{role}", f"'{who}' is not on the People tab", "")
        elif role not in [x.upper() for x in C.as_list(ppl[who].get("roles"))]:
            R.issue(cell[0], cell[1], f"answered_by.{role}", f"'{who}' does not hold the {role} role", "")
    used = set()
    for cat in K.all_catalogues():
        if cat.skill == "synthetic":
            continue
        for f in cat.fields:
            for p, _ in K.expand(f.path, P.doc):
                if not C.is_blank(C.get_path(P.doc, p)):
                    used.add(f.owner)
    for role in sorted(used | {r for _, r in P.pending}):
        if role not in P.answered_by:
            cell = P.cells.get(f"answered_by.{role}", (S_PROJECT, ""))
            R.issue(cell[0], cell[1], f"answered_by.{role}", f"nobody is named as having answered the {role} questions",
                    "name a person from the People tab who holds that role")


# ================================================================ applying
MANAGED = ("people", "policies", "sources", "silver", "metrics", "serve")
GATE_B_KEYS = ("date_ambiguity", "readback_confirmed")


def _owner_of(path: str) -> K.Field | None:
    return _field(path)


def apply(project: Project, P: Parsed, wb_file: Path) -> dict:
    """Write the parsed answers and their provenance. Assumes `validate` found no problems."""
    from .intake import load_provenance, save_provenance
    doc = P.doc
    old = project.document()
    oldprov = load_provenance(project)
    keep_prov: dict = {}
    if not P.gate_b:  # answers given after the first load survive a re-import of the main workbook
        for e, spec in (C.get_path(doc, "silver.entities") or {}).items():
            prev = C.get_path(old, f"silver.entities.{e}")
            if not isinstance(prev, dict) or not isinstance(spec, dict):
                continue
            for k in GATE_B_KEYS:
                if k not in spec and k in prev:
                    spec[k] = copy.deepcopy(prev[k])
                    base = f"silver.entities.{e}.{k}"
                    keep_prov.update({pk: rec for pk, rec in oldprov.items() if pk == base or pk.startswith(base + ".")})
    for ns in MANAGED:
        project.save_namespace(ns, doc.get(ns) if doc.get(ns) is not None else {})
    man = project.manifest
    man = man if isinstance(man, dict) else {}
    for k in ("name", "profile", "builder"):
        v = C.get_path(doc, f"project.{k}")
        if v is not None:
            man[k] = v
    project.save_namespace("project", man)

    sha = hashlib.sha256(Path(wb_file).read_bytes()).hexdigest()
    stamp_at = now_iso()
    prov = {k: v for k, v in oldprov.items()
            if not any(k == ns or k.startswith(ns + ".") for ns in MANAGED + ("project",))}
    prov.update(keep_prov)
    new = project.document()
    pending_paths = {p for p, _ in P.pending}
    kept, changed, star_kept = 0, 0, []
    for cat in K.all_catalogues():
        if cat.skill == "synthetic":
            continue
        for f in cat.fields:
            for path, caps in K.expand(f.path, new):
                if path in keep_prov or path in pending_paths:
                    continue
                v = C.get_path(new, path)
                if C.is_blank(v):
                    continue
                by = P.answered_by.get(f.owner, "")
                state = "answered"
                if C.is_token(v, C.NA_TOKEN) and not _choice(f, v):
                    state = "na"
                elif C.is_token(v, C.NONE_TOKEN) and not _choice(f, v):
                    state = "none"
                default = P.defaults.get(path, None)
                is_kept = default is not None and C.canonical_hash(default) == C.canonical_hash(v)
                cell = _cell_for(P, path)
                rec = {"by": by, "source": "workbook", "at": stamp_at, "state": state,
                       "value_hash": C.canonical_hash(v), "workbook": Path(wb_file).name, "workbook_sha256": sha[:16],
                       "cell": f"{cell[0]}!{cell[1]}" if cell else "", "default_kept": "yes" if is_kept else "no"}
                prov[path] = rec
                if is_kept:
                    kept += 1
                    if f.human_owned:
                        star_kept.append(path)
                else:
                    changed += 1
    for path, role in P.pending:
        prov[path] = {"state": "pending", "owner": P.answered_by.get(role, ""), "at": stamp_at, "source": "workbook",
                      "workbook": Path(wb_file).name}
    save_provenance(project, prov)
    keep_dir = project.loc("workbook_archive")
    keep_dir.mkdir(parents=True, exist_ok=True)
    stamp = stamp_at.replace(":", "").replace("-", "")[:15]
    shutil.copyfile(wb_file, keep_dir / f"{stamp}-{Path(wb_file).name}")
    audit(project, "intake.import", workbook=Path(wb_file).name, sha256=sha[:16], kept=kept, changed=changed,
          star_kept=len(star_kept), pending=len(P.pending))
    return {"kept": kept, "changed": changed, "star_kept": star_kept, "pending": [p for p, _ in P.pending]}


def _choice(f: K.Field, v: Any) -> bool:
    return f.type == "enum" and isinstance(v, str) and v.strip().lower() in {c.lower() for c in f.choices}


def _cell_for(P: Parsed, path: str) -> tuple | None:
    if path in P.cells:
        return P.cells[path]
    parts = C.split_path(path)
    while parts:
        p = ".".join(parts)
        if p in P.cells:
            return P.cells[p]
        if p in P.anchors:
            return P.anchors[p]
        parts = parts[:-1]
    return None


# ================================================================ validating
def validate(project: Project, P: Parsed, wb_file: Path) -> tuple[list[Issue], list[str]]:
    """Apply to a scratch copy of the project and run every intake gate there."""
    from .intake import check
    issues = list(P.issues)
    warnings: list[str] = []
    role_problem = any(i.path.startswith("answered_by.") for i in issues)
    broken_tabs = {i.sheet for i in issues if i.problem.startswith(("column '", "header row not found", "tab '"))}
    tmp = Path(tempfile.mkdtemp(prefix="dwh_import_"))
    try:
        for src in [project.root / "dwh-project.yaml", *project.spec_files()]:
            if src.exists():
                rel = src.relative_to(project.root)
                (tmp / rel).parent.mkdir(parents=True, exist_ok=True)
                shutil.copyfile(src, tmp / rel)
        (tmp / ".dwh").mkdir(exist_ok=True)
        for rel in (".dwh/profile.json", ".dwh/salt"):
            if (project.root / rel).exists():
                shutil.copyfile(project.root / rel, tmp / rel)
        rb = project.loc("readback_state")
        if rb.exists():
            shutil.copytree(rb, tmp / rb.relative_to(project.root), dirs_exist_ok=True)
        scratch = Project(tmp)
        try:
            apply(scratch, P, wb_file)
        except Exception as e:  # a shape the specs cannot hold: report it, never half-record it
            issues.append(Issue("", "", "", f"the answers could not be assembled: {type(e).__name__}: {e}", ""))
            return issues, warnings
        # with the After-first-load tab the silver gate-B questions are checked too (the data-dependent
        # parts need the warehouse, which the scratch copy does not have: they run again at build time)
        skills = [("project", "A"), ("bronze", "A"), ("silver", "B" if P.gate_b else "A"), ("gold", "A"), ("serve", "A")]
        for skill, g in skills:
            try:
                rep = check(scratch, skill, g)
            except Exception as e:
                issues.append(Issue("", "", skill, f"the {skill} answers could not be checked: {type(e).__name__}: {e}", ""))
                continue
            for e in rep.errors:
                if e.problem == "no bronze data yet":
                    continue
                if role_problem and (e.problem.startswith("recorded by '") or "has no recorded owner" in e.problem):
                    continue  # follows from the missing answerer, already reported on the Project tab
                cell = _cell_for(P, e.path) or ("", "")
                if cell[0] in broken_tabs and e.problem == "missing":
                    continue  # follows from a deleted/renamed column, already reported
                issues.append(Issue(cell[0], cell[1], e.path, e.problem, e.fix))
            for wv in rep.warnings:
                warnings.append(f"{wv.path}: {wv.problem}")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    # a dedupe: the same problem reported by two gates
    seen, out = set(), []
    for i in issues:
        k = (i.sheet, i.cell, i.path, i.problem)
        if k not in seen:
            seen.add(k)
            out.append(i)
    return out, warnings


def write_issues(src: Path, issues: list[Issue], out: Path) -> Path:
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.load_workbook(src)
    if S_ISSUES in wb.sheetnames:
        del wb[S_ISSUES]
    for ws in wb.worksheets:  # clear marks from an earlier round
        for row in ws.iter_rows():
            for c in row:
                if c.comment is not None and str(c.comment.author or "") == "dwh import":
                    c.comment = None
    ws = wb.create_sheet(S_ISSUES, 0)
    ws.sheet_view.showGridLines = False
    ws["A1"] = f"{len(issues)} problem(s) to fix — nothing was recorded yet"
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color="C00000")
    ws["A2"] = "Each problem is also marked in red on its cell (hover for the note). Fix them and send the workbook back."
    ws["A2"].font = Font(name="Arial", size=10, color=GREY_TXT)
    hdr = ["#", "Tab", "Cell", "Field", "Problem", "How to fix"]
    for i, h in enumerate(hdr, start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color=FILL_HDR, end_color=FILL_HDR)
    bad = PatternFill("solid", start_color=FILL_BAD, end_color=FILL_BAD)
    for n, iss in enumerate(issues, start=1):
        vals = [n, iss.sheet, iss.cell, iss.path, iss.problem, iss.fix]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=4 + n, column=i, value=v)
            c.font = Font(name="Arial", size=10)
        if iss.sheet in wb.sheetnames and iss.cell:
            if iss.sheet in wb.sheetnames:
                ws.cell(row=4 + n, column=2).hyperlink = f"#'{iss.sheet}'!{iss.cell}"
            target = wb[iss.sheet][iss.cell]
            target.fill = bad
            text = iss.problem + (f"\n→ {iss.fix}" if iss.fix else "")
            if target.comment is not None and str(target.comment.author) == "dwh import":
                text = target.comment.text + "\n" + text
            target.comment = Comment(text[:1000], "dwh import")
    for col, wd in zip("ABCDEF", (5, 18, 8, 46, 60, 50)):
        ws.column_dimensions[col].width = wd
    wb.active = 0
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def star_kept_summary(project: Project, paths: list[str]) -> list[str]:
    """Group ★ answers kept at their workbook default by question, with the values chosen."""
    doc = project.document()
    groups: dict[str, list[str]] = {}
    for p in paths:
        f = _field(p)
        groups.setdefault(f.path if f else p, []).append(p)
    out = []
    for pat, ps in groups.items():
        vals: dict[str, int] = {}
        for p in ps:
            raw = C.get_path(doc, p)
            v = (f"{len(raw)} rule(s): {', '.join(raw)}" if isinstance(raw, dict) else _fmt(raw))
            vals[v] = vals.get(v, 0) + 1
        shown = ", ".join(f"{v} ×{n}" if len(ps) > 1 else v for v, n in sorted(vals.items(), key=lambda x: -x[1]))
        label = pat if len(ps) > 1 else ps[0]
        out.append(f"`{label}`" + (f" ({len(ps)})" if len(ps) > 1 else "") + f": {shown}")
    return out


def star_kept_paths(project: Project) -> list[str]:
    from .intake import load_provenance
    out = []
    for p, rec in load_provenance(project).items():
        if isinstance(rec, dict) and rec.get("default_kept") == "yes" and rec.get("source") == "workbook":
            f = _field(p)
            if f is not None and f.human_owned:
                out.append(p)
    return sorted(out)


def import_workbook(project: Project, path: str, dry_run: bool = False, out: str | None = None) -> tuple[int, str]:
    src = Path(path) if Path(path).is_absolute() else Path.cwd() / path
    if not src.exists():
        src = project.root / path
    if not src.exists():
        raise WorkbookError(f"workbook not found: {path}")
    P = parse(src)
    issues, warnings = validate(project, P, src)
    lines = []
    if issues:
        dest = Path(out) if out else src.with_name(src.stem.replace("-issues", "") + "-issues.xlsx")
        write_issues(src, issues, dest)
        lines.append(f"## Workbook import — BLOCKED: {len(issues)} problem(s); nothing was recorded")
        lines.append(f"Issues workbook: {dest}")
        by_sheet: dict[str, list[Issue]] = {}
        for i in issues:
            by_sheet.setdefault(i.sheet or "(general)", []).append(i)
        for sheet, items in by_sheet.items():
            lines.append(f"\n**{sheet}**")
            for i in items[:60]:
                fix = f" → {i.fix}" if i.fix else ""
                lines.append(f"- {i.cell or '—'} `{i.path}`: {i.problem}{fix}")
            if len(items) > 60:
                lines.append(f"- … and {len(items) - 60} more on this tab")
        audit(project, "intake.import.blocked", workbook=src.name, issues=len(issues))
        return 2, "\n".join(lines)
    if dry_run:
        return 0, f"## Workbook import — dry run: no problems found ({len(P.cells)} answer cells read). Nothing recorded."
    res = apply(project, P, src)
    lines.append("## Workbook import — recorded")
    lines.append(f"{res['kept'] + res['changed']} answers recorded from {src.name}: {res['changed']} changed from the "
                 f"default, {res['kept']} defaults kept.")
    who = ", ".join(f"{r} = {p}" for r, p in sorted(P.answered_by.items()))
    lines.append(f"Recorded in the names of: {who}.")
    if res["star_kept"]:
        lines.append(f"\n★ owner decisions taken by keeping the proposed default ({len(res['star_kept'])}) — they are "
                     "the owner's answers now; `dwh status` lists them too:")
        lines += [f"- {t}" for t in star_kept_summary(project, res["star_kept"])]
    if res["pending"]:
        lines.append(f"\nPending ★ ({len(res['pending'])}) — the build may continue; release to consumers stays blocked:")
        for p in res["pending"]:
            lines.append(f"- `{p}`")
    if warnings:
        lines.append("\nWarnings:")
        lines += [f"- {w}" for w in warnings[:30]]
    lines.append("\nNext: `dwh intake check all`, then `dwh build bronze`.")
    return 0, "\n".join(lines)
