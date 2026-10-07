"""The Excel intake workbook (dwh-init 1.1): analyse → draft → workbook → (person fills) → import.

Proves: every layer's questions are in the workbook with defaults; nothing but counts leaves the
data; an unfilled or inconsistent workbook records NOTHING and comes back with every problem on its
cell; a clean one is recorded in the owners' names with cell + default-kept provenance; ★ defaults
kept are listed; the after-first-load round (read-back has no default); removed items are not
resurrected; lookups on hashed keys match (silver fix); the guarded kernel upgrade.
"""
from __future__ import annotations

import csv
import re
import hashlib
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import yaml

sys.path.insert(0, str(Path(__file__).parent))
from common import INSTALL, SKILLS, Fail, check, duck, dwh, install, lay, loc, provenance  # noqa: E402

SECRET_EMAIL = "SECRET_VALUE_123@example.com"
SECRET_STATE = "ZZSECRET"
HDR = 4


# ---------------------------------------------------------------- simulated person
def qa(ws, qid, value):
    for r in range(6, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == qid:
            ws.cell(row=r, column=7).value = value
            return
    raise Fail(f"question {qid} not on tab {ws.title}")


def col_of(ws, label):
    for c in range(1, ws.max_column + 1):
        if str(ws.cell(row=HDR, column=c).value or "").replace(" ★", "") == label:
            return c
    raise Fail(f"column {label} not on tab {ws.title}")


def grid_set(ws, match: dict, label, value):
    cols = {k: col_of(ws, k) for k in match}
    tgt, hit = col_of(ws, label), 0
    for r in range(6, ws.max_row + 1):
        if all(str(ws.cell(row=r, column=cols[k]).value or "") == v for k, v in match.items()):
            ws.cell(row=r, column=tgt).value = value
            hit += 1
    if not hit:
        raise Fail(f"{match} not on tab {ws.title}")


def grid_add(ws, values: dict):
    r = 6
    while any(ws.cell(row=r, column=c).value not in (None, "") for c in range(1, ws.max_column + 1)):
        r += 1
    for label, v in values.items():
        ws.cell(row=r, column=col_of(ws, label)).value = v


def kpi_col(ws, name):
    for c in range(6, ws.max_column + 1):
        if ws.cell(row=HDR, column=c).value == name:
            return c
    raise Fail(f"KPI {name} not on the KPIs tab")


def all_text(path: Path) -> str:
    wb = openpyxl.load_workbook(path)
    return "\n".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None)


def spec_hash(root: Path) -> str:
    h = hashlib.sha256()
    for p in [root / "dwh-project.yaml", *lay(root).spec_files()]:
        rel = p.relative_to(root).as_posix()
        h.update(rel.encode() + (p.read_bytes() if p.exists() else b"-"))
    return h.hexdigest()


# ---------------------------------------------------------------- data + draft
def write_data(root: Path) -> None:
    raw = root / "data" / "raw"
    for b, rows in (("01", [("o1", "c1", "delivered", "2024-01-03 10:00:00", "12.50", "1234"),
                           ("o2", "c2", "delivered", "2024-01-03 11:30:00", "0", "01234"),
                           ("o3", "c3", "canceled", "2024-01-04 09:15:00", "40.00", "54321"),
                           ("o4", "c1", "delivered", "2024-01-04 18:45:00", "-5.00", "1234")]),
                    ("02", [("o5", "c2", "delivered", "2024-02-01 08:00:00", "20.00", "01234"),
                           ("o6", "c4", "shipped", "2024-02-02 12:00:00", "33.10", "777"),
                           ("o7", "c3", "delivered", "2024-02-02 13:00:00", "8.90", "54321")])):
        d = raw / f"batch_{b}"
        d.mkdir(parents=True, exist_ok=True)
        with open(d / "orders.csv", "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(["order_id", "customer_id", "status", "ordered_at", "amount", "zip"])
            w.writerows(rows)
        custs = {"01": [("c1", SECRET_EMAIL, "SP"), ("c2", "b@example.com", SECRET_STATE), ("c3", "c@example.com", "RJ")],
                 "02": [("c4", "d@example.com", "SP")]}[b]
        with open(d / "customers.csv", "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(["customer_id", "email", "state"])
            w.writerows(custs)
    with open(raw / "regions.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["state", "region"])
        w.writerows([("SP", "Southeast"), ("RJ", "Southeast"), (SECRET_STATE, "North")])


DRAFT = {
    "people": {"ana": {"name": "Ana Silva", "roles": ["DE", "PO", "SME", "GOV"]}},
    "policies": {"timezone": {"reporting": "America/Sao_Paulo"}},
    "sources": {
        "orders": {"location": "data/raw/batch_{batch}/orders.csv",
                   "schema": {"columns": {"order_id": {"key": "yes"}}}},
        "customers": {"location": "data/raw/batch_{batch}/customers.csv",
                      "schema": {"columns": {"customer_id": {"key": "yes"}}}},
        "regions": {"location": "data/raw/regions.csv", "load_type": "full_refresh", "role": "reference",
                    "schema": {"columns": {"state": {"key": "yes"}}}},
    },
    "silver": {"entities": {
        "orders": {"sources": ["orders"],
                   "mapping": {"orders": {"zip": "lpad(zip, 5, '0')"}},
                   "hard_rejects": {"negative_amount": {"predicate": "amount < 0", "reason": "negative_amount"}},
                   "valid_anomalies": {"free_order": {"predicate": "amount = 0", "reason": "vouchers make real 0 orders"}},
                   "lookups": {"cust": {"reference": "customers", "on": {"customer_id": "customer_id"},
                                        "columns": ["state"], "missing": "reject"}}},
        "customers": {"sources": ["customers"]},
        "regions": {"sources": ["regions"]},
    }},
    "metrics": {
        "orders_placed": {"entity": "orders", "plain_definition": "Orders per day and customer state.",
                          "metric_type": "count", "numerator": {"measure": "count", "filter": "none"},
                          "filters": "none", "grain": ["ordered_at", "cust_state"], "time_grain": "day",
                          "date_basis": "ordered_at", "unit": "count"},
        "revenue": {"entity": "orders", "plain_definition": "Order amounts per day, canceled excluded.",
                    "metric_type": "sum", "numerator": {"measure": "sum:amount", "filter": "none"},
                    "filters": "status <> 'canceled'", "grain": ["ordered_at"], "time_grain": "day",
                    "date_basis": "ordered_at", "unit": "currency:BRL",
                    "reconciliation": {"control": "sum:amount", "excluded_by": "status = 'canceled'",
                                       "max_excluded_pct": "60"}},
    },
    "serve": {"title": "Shop (test)", "audience": "validation",
              "kpis": {"orders": {"metric": "orders_placed", "label": "Orders"}, "rev": {"metric": "revenue", "label": "Revenue"}},
              "charts": {"rev_trend": {"type": "line", "metric": "revenue", "x": "ordered_at", "color": "NA", "title": "NA"}}},
    "notes": {"silver.entities.orders.hard_rejects.negative_amount": "test: one negative order planted"},
}


def run(root: Path, log: list) -> dict:
    install(root)
    write_data(root)
    # ---------------- W1 analyse: batch folders grouped into one source each, counts only
    rc, out = dwh(root, "intake", "analyze", "data/raw")
    check("data/raw/batch_{batch}/orders.csv" in out and "data/raw/batch_{batch}/customers.csv" in out
          and "data/raw/regions.csv" in out,
          "W1 analyze groups batch_01/batch_02 files into one source with a {batch} location", log)
    check(SECRET_EMAIL not in out and SECRET_STATE not in out,
          "W1 analysis output carries counts only — no data value appears", log)
    draft = loc(root, "draft")
    check(draft.exists(), "W1 analyze starts the draft (config/intake/draft.yaml or intake/draft.yaml) with the proposed sources", log)
    draft.write_text(yaml.safe_dump(DRAFT, sort_keys=False), encoding="utf-8")

    # ---------------- W2 workbook: every layer, defaults, dropdowns, no values
    rc, out = dwh(root, "intake", "workbook")
    wbp = loc(root, "workbook_out") / "shop-intake.xlsx"
    if not wbp.exists():
        wbp = next(loc(root, "workbook_out").glob("*-intake.xlsx"))
    wb = openpyxl.load_workbook(wbp)
    want = ["Start here", "People", "Project", "Governance", "Reporting policy", "Sources", "Schema", "Classification",
            "Silver model", "Entities", "Data quality", "Lookups & flags", "KPIs", "Golden values", "Dashboard",
            "Dashboard visuals"]
    check(all(t in wb.sheetnames for t in want), "W2 one tab per kind of question, every layer (16 tabs)", log)
    meta = wb["_meta"]
    defaults = {meta.cell(row=r, column=1).value for r in range(9, meta.max_row + 1)}
    from_catalogues = ["policies.compliance", "policies.egress", "policies.timezone.reporting", "policies.money.scale",
                       "sources.orders.load_type", "sources.orders.schema.columns.amount.type",
                       "sources.customers.schema.columns.email.classification", "silver.entities.orders.natural_key",
                       "silver.entities.orders.merge.strategy", "silver.entities.orders.null_policy.amount",
                       "silver.entities.customers.masking.email", "metrics.revenue.metric_type",
                       "metrics.revenue.grain", "serve.reference_clock", "serve.row_security"]
    check(all(p in defaults for p in from_catalogues),
          "W2 project, governance, bronze, silver, gold and serve questions all carry a default", log)
    vals = {}
    ws = wb["Schema"]
    for r in range(6, ws.max_row + 1):
        if ws.cell(row=r, column=2).value == "amount":
            vals["amount"] = ws.cell(row=r, column=col_of(ws, "Type")).value
        if ws.cell(row=r, column=2).value == "zip":
            vals["zip"] = ws.cell(row=r, column=col_of(ws, "Type")).value
    check(vals.get("amount", "").startswith("DECIMAL") and vals.get("zip") == "VARCHAR",
          "W2 analysis defaults: money → DECIMAL, a ragged code (zip) → text", log)
    ws = wb["Classification"]
    cls = {(ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value): ws.cell(row=r, column=3).value
           for r in range(6, ws.max_row + 1)}
    check(cls.get(("customers", "email")) == "pii" and cls.get(("orders", "status")) == "internal",
          "W2 classification defaults from column names (email → pii)", log)
    check(len(wb["Sources"].data_validations.dataValidation) > 0 and len(wb["KPIs"].data_validations.dataValidation) > 0,
          "W2 answer cells with a fixed list carry a drop-down", log)
    g = wb["Golden values"]
    gv = [g.cell(row=r, column=3).value for r in range(6, g.max_row + 1) if g.cell(row=r, column=1).value in DRAFT["metrics"]]
    check(gv and all(v in (None, "") for v in gv), "W2 golden values have no default (blank rows per KPI)", log)
    text = all_text(wbp)
    check(SECRET_EMAIL not in text and SECRET_STATE not in text,
          "W2 egress: the workbook holds no data value (counts and names only)", log)

    # ---------------- W3 unfilled people → blocked, NOTHING recorded, issues workbook
    empty = root / "intake" / "unfilled.xlsx"
    wb2 = openpyxl.load_workbook(wbp)
    qa(wb2["Project"], "answered_by.GOV", "")
    wb2["Golden values"]  # untouched
    grid_set(wb2["Schema"], {"Source": "orders", "Column": "amount"}, "Type", "NA")
    wb2.save(empty)
    before = spec_hash(root)
    rc, out = dwh(root, "intake", "import", str(empty), expect=2)
    check(spec_hash(root) == before, "W3 a refused import records nothing (specs and provenance unchanged)", log)
    iss = root / "intake" / "unfilled-issues.xlsx"
    check(iss.exists() and "Issues" in openpyxl.load_workbook(iss).sheetnames,
          "W3 an -issues.xlsx copy is written with an Issues tab", log)
    ib = openpyxl.load_workbook(iss)
    rows = [[ib["Issues"].cell(row=r, column=c).value for c in range(1, 6)] for r in range(5, ib["Issues"].max_row + 1)]
    hit_schema = [r for r in rows if r[1] == "Schema" and "amount.type" in str(r[3]) and "NA is not allowed" in str(r[4])]
    hit_role = [r for r in rows if r[1] == "Project" and "answered_by.GOV" in str(r[3])]
    check(bool(hit_schema) and bool(hit_role), "W3 each problem names its tab, cell and fix (schema NA, missing GOV answerer)", log)
    cell = ib["Schema"][hit_schema[0][2]]
    check(cell.comment is not None and cell.fill.start_color.rgb.endswith("FFC7CE"),
          "W3 the problem cell is marked red with a note", log)

    # ---------------- W4 errors are caught on their cells
    def mutate(name, fn, expect_text):
        w = openpyxl.load_workbook(wbp)
        fn(w)
        p = root / "intake" / f"err_{name}.xlsx"
        w.save(p)
        rc, out = dwh(root, "intake", "import", str(p), "--dry-run", expect=None)
        return rc == 2 and all(t in out for t in expect_text)
    ok = [
        mutate("enum", lambda w: grid_set(w["Sources"], {"Source": "orders"}, "Load type", "incremental"),
               ["Sources", "load_type", "must be one of"]),
        mutate("sql", lambda w: grid_set(w["Data quality"], {"Entity": "orders", "Column / rule name": "negative_amount"},
                                         "Condition (SQL)", "amount < < 0"), ["Data quality", "predicate"]),
        mutate("assistant", lambda w: grid_add(w["People"], {"Person id": "claude", "Full name": "Claude", "Roles": "GOV"}),
               ["assistant cannot be listed"]),
        mutate("golden", lambda w: grid_set(w["Golden values"], {"KPI": "revenue"}, "Key (grain values)", "ordered_at 2024-01-03"),
               ["Golden values", "is not col=value"]),
        mutate("grain", lambda w: w["KPIs"].cell(row=next(r for r in range(6, 40) if w["KPIs"].cell(row=r, column=1).value == "grain"),
                                                 column=kpi_col(w["KPIs"], "revenue")).__setattr__("value", "ordered_at, region"),
               ["KPIs", "grain column 'region' not in entity"]),
        mutate("governed", lambda w: qa(w["Project"], "project.profile", "governed"), ["golden_values", "pending"]),
    ]
    check(all(ok), f"W4 bad drop-down value, bad SQL, assistant as a person, malformed golden key, unknown grain "
                   f"column, governed + pending golden → each refused on its cell ({sum(ok)}/6)", log)
    rc, out = dwh(root, "intake", "import", str(wbp), "--dry-run")
    check(spec_hash(root) == before and "no problems found" in out, "W4 --dry-run checks and records nothing", log)

    # ---------------- W5 clean import: owners, cells, default_kept
    filled = root / "intake" / "filled.xlsx"
    w = openpyxl.load_workbook(wbp)
    grid_set(w["Entities"], {"Entity": "orders"}, "Max % of a batch rejected", "40")   # a changed default
    w.save(filled)
    rc, out = dwh(root, "intake", "import", str(filled))
    prov = provenance(root)
    r_cls = prov.get("sources.customers.schema.columns.email.classification") or {}
    r_tol = prov.get("silver.entities.orders.dead_letter_tolerance_pct") or {}
    check(r_cls.get("by") == "ana" and r_cls.get("source") == "workbook" and r_cls.get("cell", "").startswith("Classification!")
          and r_cls.get("default_kept") == "yes",
          "W5 a kept ★ default is recorded as the GOV owner's answer, with its cell and default_kept: yes", log)
    check(r_tol.get("default_kept") == "no" and r_tol.get("cell", "").startswith("Entities!"),
          "W5 a changed default is recorded with default_kept: no", log)
    check(any(loc(root, "workbook_archive").glob("*filled.xlsx")),
          "W5 the imported workbook is kept as received in the workbook archive", log)
    rc, st = dwh(root, "status")
    check("decisions taken by keeping a proposed workbook default" in st and "classification" in st,
          "W5 dwh status lists the ★ decisions taken by keeping a default", log)
    rc, out = dwh(root, "intake", "check", "gold", "--json", expect=None)
    check(rc == 0 and "golden_values" in out and "pending" in out,
          "W5 KPIs without golden values are pending in the fast profile (build may go on, release blocked)", log)
    for skill in ("project", "bronze", "gold", "serve"):
        dwh(root, "intake", "check", skill)
    check(True, "W5 project, bronze, gold and serve gates pass straight from the workbook", log)

    # ---------------- W6 after the first load: read-back has no default
    dwh(root, "build", "bronze")
    rc, out = dwh(root, "intake", "check", "silver", "--gate", "B", expect=2)
    rc, out = dwh(root, "intake", "workbook", "--out", "intake/round2.xlsx")
    w = openpyxl.load_workbook(root / "intake" / "round2.xlsx")
    check("After first load" in w.sheetnames, "W6 regenerating after the bronze load adds the After first load tab", log)
    ws = w["After first load"]
    rb = [r for r in range(6, ws.max_row + 1) if ws.cell(row=r, column=2).value == "Rule read-back"]
    check(rb and all(ws.cell(row=r, column=5).value in (None, "") for r in rb)
          and all("rows" in str(ws.cell(row=r, column=4).value) for r in rb),
          "W6 the rule read-back shows match counts and has no default", log)
    rc, out = dwh(root, "intake", "import", str(root / "intake" / "round2.xlsx"), expect=2)
    check("readback_confirmed" in out, "W6 importing without confirming the read-back is refused", log)
    for r in rb:
        ws.cell(row=r, column=5).value = "yes"
    w.save(root / "intake" / "round2.xlsx")
    dwh(root, "intake", "import", "intake/round2.xlsx")
    dwh(root, "intake", "check", "silver", "--gate", "B")
    check(True, "W6 after the owner confirms the read-back in the workbook, silver gate B passes", log)
    dwh(root, "intake", "workbook", "--out", "intake/round2b.xlsx")
    rc, out = dwh(root, "intake", "check", "silver", "--gate", "B", expect=None)
    wsb = openpyxl.load_workbook(root / "intake" / "round2b.xlsx")["After first load"]
    shown = [wsb.cell(row=r, column=5).value for r in range(6, wsb.max_row + 1)
             if wsb.cell(row=r, column=2).value == "Rule read-back"]
    check(rc == 0 and shown and all(v == "yes" for v in shown),
          "W6 regenerating the workbook keeps a confirmed read-back confirmed (and shows it)", log)

    # ---------------- W7 lookups on a hashed key match (silver fix)
    dwh(root, "build", "silver")
    misses = duck(root, "SELECT COUNT(*) FROM silver_orders WHERE cust_state IS NULL")[0][0]
    n = duck(root, "SELECT COUNT(*) FROM silver_orders")[0][0]
    hashed = duck(root, "SELECT COUNT(*) FROM silver_customers WHERE length(customer_id) = 32")[0][0]
    check(hashed > 0 and n == 6 and misses == 0,
          f"W7 customer_id hashed in both entities, the lookup still matches every order ({n} kept, {misses} missing)", log)
    dwh(root, "build", "gold")

    # ---------------- W8 a removed item is not resurrected; earlier gate-B answers survive
    w = openpyxl.load_workbook(root / "intake" / "round2.xlsx")
    ws = w["KPIs"]
    c = kpi_col(ws, "orders_placed")
    ws.cell(row=HDR, column=c).value = None
    vis = w["Dashboard visuals"]
    for r in range(6, vis.max_row + 1):
        if vis.cell(row=r, column=col_of(vis, "KPI")).value == "orders_placed":
            for cc in range(1, vis.max_column + 1):
                vis.cell(row=r, column=cc).value = None
    w.save(root / "intake" / "round3.xlsx")
    dwh(root, "intake", "import", "intake/round3.xlsx")
    dwh(root, "intake", "workbook", "--out", "intake/round4.xlsx")
    names = [openpyxl.load_workbook(root / "intake" / "round4.xlsx")["KPIs"].cell(row=HDR, column=cc).value
             for cc in range(6, 12)]
    check("orders_placed" not in names and "revenue" in names,
          "W8 a KPI the owner removed is not brought back by the draft on regeneration", log)
    dwh(root, "intake", "import", str(filled))   # a workbook from BEFORE the first load
    rc, out = dwh(root, "intake", "check", "silver", "--gate", "B", expect=None)
    check(rc == 0, "W8 read-back answers given after the first load survive re-importing an older workbook", log)

    # ---------------- W9 guarded kernel upgrade
    KV = re.search(r'VERSION = "([^"]+)"', (SKILLS / "dwh-init" / "kernel" / "dwh_core" / "__init__.py").read_text()).group(1)
    init = root / "dwh_core" / "__init__.py"
    init.write_text(init.read_text().replace(f'VERSION = "{KV}"', 'VERSION = "1.0.0"'))
    r = subprocess.run([sys.executable, str(INSTALL), "--project", str(root)], capture_output=True, text=True)
    check(r.returncode == 2 and "--upgrade" in r.stdout, "W9 an older pinned kernel is not replaced silently", log)
    r = subprocess.run([sys.executable, str(INSTALL), "--project", str(root), "--upgrade"], capture_output=True, text=True)
    man = (root / "dwh-project.yaml").read_text()
    check(r.returncode == 0 and f'dwh_core_version: "{KV}"' in man and f'VERSION = "{KV}"' in init.read_text(),
          "W9 --upgrade replaces the kernel and re-pins the manifest (same key algorithm)", log)
    return {}
